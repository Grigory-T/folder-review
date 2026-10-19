"""Owned synthetic Windows portable-folder acceptance, with no private dataset access."""

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import time
from pathlib import Path

base = Path(__file__).resolve().parent.parent
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--exe", required=True, type=Path)
parser.add_argument("--results", required=True, type=Path)
args = parser.parse_args()
if os.name != "nt":
    parser.error("Run portable acceptance on Windows")
candidate = args.exe.resolve().parent
work = base / ".build/portable-acceptance"
if work.exists():
    shutil.rmtree(work)
work.mkdir()
tool = work / "portable folder-пример"
shutil.copytree(
    candidate, tool, ignore=shutil.ignore_patterns("folder-review-data", "*.xlsx")
)
source = work / "input"
source.mkdir(exist_ok=True)
for n in range(120):
    (source / f"file-{n}.txt").write_bytes((f"payload-{n} " * n).encode())
(source / "copy.txt").write_bytes((source / "file-119.txt").read_bytes())
import zipfile

with zipfile.ZipFile(source / "data.zip", "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("notes.txt", b"archive payload")
canary = work / "outside"
canary.mkdir(exist_ok=True)
env = os.environ.copy()
for key in ("TEMP", "TMP", "TMPDIR", "LOCALAPPDATA", "APPDATA", "PYTHONPYCACHEPREFIX"):
    env[key] = str(canary)


def snapshot(root):
    return {
        p.relative_to(root).as_posix(): (
            hashlib.sha256(p.read_bytes()).hexdigest(),
            p.stat().st_mtime_ns,
        )
        for p in root.rglob("*")
        if p.is_file()
    }


before = snapshot(source)
timings = []


def invoke(folder, args, expected=0):
    start = time.perf_counter()
    run = subprocess.run(
        [str(folder / "folder-review.exe"), *args],
        env=env,
        cwd=canary,
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert run.returncode == expected, (run.returncode, run.stdout, run.stderr)
    timings.append(round(time.perf_counter() - start, 3))
    assert not list(canary.iterdir()), (
        "Application wrote into the external working/temp/profile canary"
    )
    return run


for mode in ("off", "filtered", "all"):
    invoke(tool, ["--root", str(source), "--duplicates", mode])
    assert (tool / "folder-review.xlsx").is_file()
    assert list((tool / "folder-review-data/logs").glob("*.jsonl"))
    assert not list((tool / "folder-review-data").glob("work-*"))
assert snapshot(source) == before
# Every log is neutral aggregate data, never source names/paths/content.
for log in (tool / "folder-review-data/logs").glob("*.jsonl"):
    text = log.read_text()
    assert str(source) not in text and "file-119" not in text and "payload-" not in text
    assert json.loads(text.splitlines()[-1])["exit_code"] == 0
invoke(tool, ["--root", str(tool)])
# Inspect the generated report with the independent build-only reader.
from openpyxl import load_workbook

book = load_workbook(tool / "folder-review.xlsx", read_only=True)
sheet = book["folder_review"]
head = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
names = [
    r[head.index("File name")] for r in sheet.iter_rows(min_row=7, values_only=True)
]
book.close()
assert names == ["folder-review.exe"], names
report_hash = hashlib.sha256((tool / "folder-review.xlsx").read_bytes()).hexdigest()
invoke(tool, ["--root", str(source), "--output", str(canary / "escape.xlsx")], 1)
invoke(tool, ["--root", str(source), "--log", str(canary / "escape.jsonl")], 1)
assert (
    hashlib.sha256((tool / "folder-review.xlsx").read_bytes()).hexdigest()
    == report_hash
)
invoke(tool, ["--root", str(tool / "_internal")], 1)
invoke(tool, ["--root", str(tool / "folder-review-data")], 1)
# Copy the entire tool folder, retaining report and logs, and run from elsewhere.
relocated = work / "copied tool"
shutil.copytree(tool, relocated, dirs_exist_ok=True)
invoke(relocated, ["--root", str(source)])
assert snapshot(source) == before
assert (relocated / "folder-review.xlsx").is_file()
# Junction redirection is rejected before any writes through it.
support = relocated / "folder-review-data"
saved = relocated / "saved-support"
support.rename(saved)
subprocess.run(
    ["cmd.exe", "/c", "mklink", "/J", str(support), str(canary)],
    capture_output=True,
    check=True,
)
try:
    invoke(relocated, ["--root", str(source)], 1)
finally:
    support.rmdir()
    saved.rename(support)
# A write-denied tool folder fails rather than choosing another storage location.
readonly = work / "write-denied tool"
shutil.copytree(
    candidate, readonly, ignore=shutil.ignore_patterns("folder-review-data", "*.xlsx")
)
import ctypes

size = ctypes.c_uint(32768)
name = ctypes.create_unicode_buffer(size.value)
assert ctypes.windll.advapi32.GetUserNameW(name, ctypes.byref(size))
user = name.value
subprocess.run(
    ["icacls.exe", str(readonly), "/deny", user + ":(WD,AD)"],
    check=True,
    capture_output=True,
)
try:
    invoke(readonly, ["--root", str(source)], 1)
finally:
    subprocess.run(
        ["icacls.exe", str(readonly), "/remove:d", user],
        check=True,
        capture_output=True,
    )
assert not (readonly / "folder-review.xlsx").exists()
assert not list((readonly / "folder-review-data").glob("work-*"))
receipt = {
    "exe_sha256": hashlib.sha256(
        (candidate / "folder-review.exe").read_bytes()
    ).hexdigest(),
    "standard_onedir": True,
    "all_duplicate_modes": True,
    "external_temp_profile_working_directory_canary_untouched": True,
    "input_bytes_and_modification_times_unchanged": True,
    "report_and_logs_beside_exe": True,
    "temporary_workspaces_cleaned": True,
    "self_scan_excludes_support_runtime_and_report": True,
    "whole_folder_relocation": True,
    "external_report_and_log_rejected": True,
    "support_junction_rejected": True,
    "unwritable_tool_folder_fails_without_fallback": True,
    "run_seconds": timings,
}
args.results.write_text(json.dumps(receipt, indent=2))
print(json.dumps(receipt))
