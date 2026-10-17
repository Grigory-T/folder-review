"""Compare the shipped executable with the original CLI on synthetic inputs."""

from __future__ import annotations

import argparse
import binascii
import hashlib
import json
import marshal
import ntpath
import os
import struct
import subprocess
import sys
import tempfile
import tomllib
import types
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import pefile
import py7zr
from blake3 import blake3
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PyInstaller.archive.readers import CArchiveReader

ROOT = Path(__file__).resolve().parent.parent
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def normalized_code(code):
    """Ignore build-path filenames, but preserve bytecode and all constants."""
    constants = tuple(
        normalized_code(value) if isinstance(value, types.CodeType) else value
        for value in code.co_consts
    )
    return code.replace(co_filename="folder_review.py", co_consts=constants)


def inspect_payload(exe):
    with pefile.PE(str(exe)) as pe:
        info = pe.VS_FIXEDFILEINFO[0]
        version = (
            info.FileVersionMS >> 16,
            info.FileVersionMS & 65535,
            info.FileVersionLS >> 16,
        )
        expected = tuple(
            int(part)
            for part in tomllib.loads((ROOT / "pyproject.toml").read_text())["project"][
                "version"
            ].split(".")
        )
        assert version == expected, (version, expected)
    archive = CArchiveReader(str(exe))
    pyz = archive.open_embedded_archive("PYZ.pyz")
    embedded = pyz.extract("folder_review")
    original = compile(
        (ROOT / "folder_review.py").read_bytes(), "folder_review.py", "exec", optimize=0
    )
    assert marshal.dumps(normalized_code(embedded)) == marshal.dumps(
        normalized_code(original)
    ), "Embedded scanner bytecode differs from the source"
    filenames = set()

    def inspect_code(code):
        filenames.add(code.co_filename)
        assert not os.path.isabs(code.co_filename) and not ntpath.isabs(
            code.co_filename
        ), ("Absolute build path in embedded Python code", code.co_filename)
        for constant in code.co_consts:
            if isinstance(constant, types.CodeType):
                inspect_code(constant)

    for name, entry in pyz.toc.items():
        if entry[0] in (0, 1):  # modules/packages, excluding namespace entries
            inspect_code(pyz.extract(name))
    for name, entry in archive.toc.items():
        if entry[-1] == "s":
            inspect_code(marshal.loads(archive.extract(name)))
    import io

    with zipfile.ZipFile(io.BytesIO(archive.extract("base_library.zip"))) as standard:
        assert standard.namelist() == sorted(standard.namelist()), (
            "Standard-library ZIP entries must have deterministic ordering"
        )
        for name in standard.namelist():
            if name.endswith(".pyc"):
                inspect_code(marshal.loads(standard.read(name)[16:]))

    pins = json.loads((ROOT / "windows_build/7zip-pins.json").read_text())
    for name in ("7z.exe", "7z.dll"):
        assert hashlib.sha256(archive.extract(name)).hexdigest() == pins["sha256"][name]
    for module in ("setuptools", "pkg_resources", "packaging"):
        assert not any(
            name == module or name.startswith(module + ".") for name in pyz.toc
        )
    ui_code = pyz.extract("folder_review_ui")
    ui_original = compile(
        (ROOT / "folder_review_ui.py").read_bytes(),
        "folder_review_ui.py",
        "exec",
        optimize=0,
    )
    assert marshal.dumps(normalized_code(ui_code)) == marshal.dumps(
        normalized_code(ui_original)
    ), "Embedded UI differs from source"
    for module in (
        "folder_review_ui",
        "tkinter",
        "folder_review",
        "xlsxwriter",
        "blake3",
        "olefile",
        "py7zr",
        "rarfile",
    ):
        assert module in pyz.toc, module
    assert not any(
        name == "openpyxl" or name.startswith("openpyxl.") for name in pyz.toc
    )
    return {
        "scanner_bytecode_identical": True,
        "embedded_code_paths_relative": True,
        "distinct_code_filenames_checked": len(filenames),
        "bundled_7zip_hashes_verified": True,
        "builder_packages_excluded": True,
    }


def make_zip(path, members, compression=zipfile.ZIP_DEFLATED):
    with zipfile.ZipFile(path, "w", compression) as archive:
        for name, value in members.items():
            archive.writestr(
                name, value.read_bytes() if isinstance(value, Path) else value
            )
    return path


def make_seven(path, members, **kwargs):
    with py7zr.SevenZipFile(path, "w", **kwargs) as archive:
        for name, value in members.items():
            if isinstance(value, Path):
                archive.write(value, name)
            else:
                archive.writestr(value, name)
    return path


def make_unusable_seven(path):
    # py7zr correctly refuses to write unusable member paths. Create one
    # stored member, then change the unencoded UTF-16 name to '.' and repair
    # the two header CRCs. The payload and its integrity hash stay intact.
    with py7zr.SevenZipFile(path, "w", filters=[{"id": py7zr.FILTER_COPY}]) as archive:
        archive.set_encoded_header_mode(False)
        archive.writestr(b"member", "A")
    data = bytearray(path.read_bytes())
    offset, size = struct.unpack_from("<QQ", data, 12)
    start = 32 + offset
    name = data.index(b"A\x00\x00\x00", start)
    data[name] = ord(".")
    struct.pack_into("<I", data, 28, binascii.crc32(data[start : start + size]))
    struct.pack_into("<I", data, 8, binascii.crc32(data[12:32]))
    path.write_bytes(data)


def normalized_workbook(path):
    """Compare every ZIP member, cell type, style, formula and layout.

    The only omissions are access-time values, scan start/end timestamps,
    and workbook creation/modification timestamps. Independent scans may
    legitimately produce different values for these fields.
    """
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        headers = list(
            next(
                workbook["folder_review"].iter_rows(
                    min_row=6, max_row=6, values_only=True
                )
            )
        )
        access_column = get_column_letter(headers.index("Accessed at") + 1)
        status = list(workbook["scan_status"].iter_rows(values_only=True))
        volatile_status_rows = {
            index
            for index, row in enumerate(status, 1)
            if row and row[0] in ("Started at", "Completed at")
        }
        status_values = {row[0]: row[1] for row in status[1:17]}
        rows = sum(1 for _ in workbook["folder_review"].iter_rows(min_row=7))
    finally:
        workbook.close()
    result = {}
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            data = archive.read(name)
            if name in (
                "xl/worksheets/sheet1.xml",
                "xl/worksheets/sheet2.xml",
                "docProps/core.xml",
            ):
                tree = ET.fromstring(data)
                if name == "docProps/core.xml":
                    for element in tree:
                        if element.tag.rsplit("}", 1)[-1] in ("created", "modified"):
                            element.text = "VOLATILE"
                else:
                    for row in tree.findall("s:sheetData/s:row", NS):
                        number = int(row.attrib["r"])
                        for cell in row.findall("s:c", NS):
                            if (
                                name.endswith("sheet1.xml")
                                and number >= 7
                                and cell.attrib["r"] == f"{access_column}{number}"
                            ) or (
                                name.endswith("sheet2.xml")
                                and number in volatile_status_rows
                                and cell.attrib["r"] == f"B{number}"
                            ):
                                for value in cell:
                                    value.text = "VOLATILE"
                data = ET.tostring(tree)
            result[name] = data
    return result, status_values, rows


def snapshot(root):
    result = {}

    def record(path):
        metadata = path.lstat()
        value = {
            "mtime_ns": metadata.st_mtime_ns,
            "mode": metadata.st_mode,
            "attributes": getattr(metadata, "st_file_attributes", None),
            "reparse_tag": getattr(metadata, "st_reparse_tag", None),
        }
        if path.is_symlink() or path.is_junction():
            value["link_target"] = os.readlink(path)
        elif path.is_file():
            value["bytes"] = metadata.st_size
            value["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        result[str(path.relative_to(root))] = value

    record(root)
    for current, directories, files in os.walk(root, followlinks=False):
        for name in directories + files:
            record(Path(current) / name)
        directories[:] = [
            name
            for name in directories
            if not (Path(current) / name).is_junction()
            and not (Path(current) / name).is_symlink()
        ]
    return result


def main():
    if sys.flags.optimize:
        raise RuntimeError("Release verification requires Python optimization level 0")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--exe", required=True, type=Path)
    parser.add_argument("--results", required=True, type=Path)
    args = parser.parse_args()
    if os.name != "nt":
        parser.error("Run this verification on Windows")
    exe = args.exe.resolve()
    results = {
        "payload": inspect_payload(exe),
        "cases": [],
        "volatile_fields": [
            "access time",
            "scan start/end",
            "workbook creation/modification time",
        ],
    }
    minimal_env = {
        key: value
        for key, value in os.environ.items()
        if key.upper()
        not in ("PATH", "PROGRAMFILES", "PROGRAMFILES(X86)", "PYTHONPATH", "PYTHONHOME")
    }
    minimal_env["PATH"] = str(Path(os.environ["SystemRoot"]) / "System32")
    with tempfile.TemporaryDirectory(prefix="folder-review-acceptance-") as temporary:
        base = Path(temporary)
        root = base / "input"
        staging = base / "staging"
        output = base / "output"
        for folder in (root, staging, output):
            folder.mkdir()
        minimal_env["ProgramFiles"] = str(base / "absent")
        minimal_env["ProgramFiles(x86)"] = str(base / "absent")
        minimal_env["TEMP"] = str(staging)
        minimal_env["TMP"] = str(staging)

        office = make_zip(
            staging / "office.xlsx",
            {
                "[Content_Types].xml": b"<Types/>",
                "_rels/.rels": b"<Relationships/>",
                "xl/workbook.xml": b"<workbook/>",
            },
        )
        inner = make_zip(
            staging / "inner.zip",
            {
                "deep/document.pdf": b"%PDF-1.4\n",
                "deep/book.xlsx": office,
                "deep/bom.txt": b"\xff\xfeh\x00i\x00",
            },
        )
        middle = make_seven(staging / "middle.7z", {"layer/inner.zip": inner})
        make_zip(root / "nested.zip", {"outer/middle.7z": middle})
        make_seven(root / "reverse.7z", {"inner.zip": inner})
        for name, method in (
            ("zip-bzip2", zipfile.ZIP_BZIP2),
            ("zip-lzma", zipfile.ZIP_LZMA),
        ):
            make_zip(
                root / f"{name}.zip",
                {"member.txt": b"ZIP codec payload" * 3000},
                method,
            )
        make_zip(
            root / "rar-wrapped.zip",
            {"inside.rar": ROOT / "tests/fixtures/rar-with-zip.rar"},
        )
        make_zip(root / "empty.zip", {})
        make_seven(root / "empty.7z", {})
        (root / "renamed.bin").write_bytes(office.read_bytes())
        (root / "broken.zip").write_bytes(b"PK\x03\x04bad")
        (root / "~$owner.xlsx").write_bytes(b"owner")
        (root / "desktop.ini").write_text("[.ShellClassInfo]", encoding="utf-8")
        (root / "=formula-like.txt").write_text("text", encoding="utf-8")
        (root / "unicode-кириллица-文件.txt").write_text(
            "Unicode text", encoding="utf-8"
        )
        (root / "empty").write_bytes(b"")
        make_seven(
            root / "encrypted.7z", {"secret.txt": b"secret"}, password="test-password"
        )
        make_seven(
            root / "encrypted-headers.7z",
            {"secret.txt": b"secret"},
            password="test-password",
            header_encryption=True,
        )
        make_zip(root / "unusable-path.zip", {"../..": b"member"})
        make_unusable_seven(root / "unusable-path.7z")
        make_zip(root / "crc.zip", {"crc.txt": b"CRC payload"}, zipfile.ZIP_STORED)
        data = bytearray((root / "crc.zip").read_bytes())
        name_length, extra_length = struct.unpack_from("<HH", data, 26)
        data[30 + name_length + extra_length] ^= 1
        (root / "crc.zip").write_bytes(data)
        with py7zr.SevenZipFile(
            root / "crc.7z", "w", filters=[{"id": py7zr.FILTER_COPY}]
        ) as archive:
            archive.set_encoded_header_mode(False)
            archive.writestr(b"A" * 65536, "crc.txt")
        data = bytearray((root / "crc.7z").read_bytes())
        assert data[32] == ord("A")
        data[32] = ord("B")
        (root / "crc.7z").write_bytes(data)
        boms = {
            "utf8": b"\xef\xbb\xbf",
            "utf16le": b"\xff\xfe",
            "utf16be": b"\xfe\xff",
            "utf32le": b"\xff\xfe\x00\x00",
            "utf32be": b"\x00\x00\xfe\xff",
        }
        for name, bom in boms.items():
            (root / f"{name}.txt").write_bytes(bom + b"text")
        target = root / "hardlink.txt"
        target.write_text("hardlink", encoding="utf-8")
        os.link(target, root / "hardlink-copy.txt")
        junction = root / "junction"
        subprocess.run(
            ["cmd.exe", "/d", "/c", "mklink", "/J", str(junction), str(staging)],
            check=True,
            capture_output=True,
        )
        subprocess.run(
            ["attrib.exe", "+H", "+R", str(root / "desktop.ini")], check=True
        )

        deep = root
        while len(str(deep)) < 330:
            deep /= "long-component-abcdefghijk"
        deep.mkdir(parents=True)
        (deep / "archive.7z").write_bytes(middle.read_bytes())
        long_output = output
        while len(str(long_output)) < 330:
            long_output /= "long-output-component-abcdef"
        long_output.mkdir(parents=True)

        complete = base / "complete"
        complete.mkdir()
        make_zip(
            complete / "archive.zip",
            {
                "middle.7z": middle,
                "nested.rar": ROOT / "tests/fixtures/rar-with-zip.rar",
            },
        )
        rar_cases = []
        # Neutral, committed fixtures exercise RAR5 solid compression, BLAKE2,
        # encrypted headers/content, CRC failures, and multi-volume behavior.
        for fixture in sorted((ROOT / "tests/fixtures/windows").glob("*")):
            if fixture.name == "README.txt":
                continue
            if fixture.suffix.lower() != ".rar":
                import shutil

                shutil.copyfile(fixture, root / fixture.name)
                continue
            if fixture.name.startswith("rar5-split."):
                folder = base / "rar5-split"
                expected = 2
            else:
                folder = base / fixture.stem
                expected = (
                    2 if "encrypted" in fixture.stem or "corrupt" in fixture.stem else 0
                )
            folder.mkdir(exist_ok=True)
            import shutil

            shutil.copyfile(fixture, folder / fixture.name)
            if not any(case[0] == folder.name for case in rar_cases):
                rar_cases.append((folder.name, folder, [], expected, output))
        codecs = base / "codecs"
        codecs.mkdir()
        configurations = {
            name: [{"id": getattr(py7zr, "FILTER_" + name)}]
            for name in (
                "COPY",
                "LZMA",
                "LZMA2",
                "BZIP2",
                "DEFLATE",
                "PPMD",
                "ZSTD",
                "BROTLI",
            )
        }
        for name in ("X86", "ARM", "ARMTHUMB", "POWERPC", "SPARC", "IA64", "DELTA"):
            configurations[name + "-LZMA2"] = [
                {"id": getattr(py7zr, "FILTER_" + name)},
                {"id": py7zr.FILTER_LZMA2},
            ]
        for name, filters in configurations.items():
            make_seven(
                codecs / f"{name}.7z",
                {"member.txt": b"codec payload " * 3000 + bytes(range(256))},
                filters=filters,
            )

        # Keep format policy checks in the actual production EXE comparisons.
        # Arbitrary/unrecognized suffixes are included to prove expansion does
        # not depend on a growing list of document exceptions.
        leaf_input = base / "package-policy"
        leaf_input.mkdir()
        leaf_zip = make_zip(
            staging / "diagram-package.zip",
            {
                "[Content_Types].xml": b"<Types/>",
                "_rels/.rels": b"<Relationships/>",
                "visio/document.xml": b"<VisioDocument/>",
                "visio/masters/_rels/master43.xml.rels": b"<Relationships/>",
                "visio/media/image2.emf": bytes(100),
                "internal/should-not-appear.zip": b"PK\x03\x04bad",
            },
        )
        # Load only the policy constants from the same validated source.
        import importlib.util

        spec = importlib.util.spec_from_file_location(
            "release_policy", ROOT / "folder_review.py"
        )
        policy = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = policy
        spec.loader.exec_module(policy)
        leaf_names = {
            f"document.{extension.upper()}" for extension in policy.ZIP_LEAF_FORMATS
        }
        leaf_names.add("document.UNKNOWN-FORMAT")
        for name in leaf_names:
            (leaf_input / name).write_bytes(leaf_zip.read_bytes())
        leaf_members = {name: leaf_input / name for name in leaf_names}
        make_zip(leaf_input / "outer.zip", leaf_members)
        make_seven(leaf_input / "outer.7z", leaf_members)
        for name in ("leaf-packages.rar", "leaf-packages-compressed.rar"):
            (leaf_input / name).write_bytes(
                (ROOT / "tests/fixtures/windows" / name).read_bytes()
            )
        expected_leaf_rows = len(leaf_names) * 3 + 77 * 2

        hash_input = base / "hash-input"
        (hash_input / "empty" / "deep").mkdir(parents=True)
        hash_payload = b"same complete content" + bytes(range(256)) * 4096
        hash_names = {"one.txt", "two.bin", "1.2.3-cp313-cp313-win_amd64"}
        for name in hash_names:
            (hash_input / name).write_bytes(hash_payload)
        (hash_input / "empty-file").write_bytes(b"")
        make_zip(
            hash_input / "duplicates.zip",
            {"inside.txt": hash_payload, "empty/deep/": b""},
        )

        marker_input = base / "marker-input"
        (marker_input / "empty").mkdir(parents=True)
        make_seven(
            marker_input / "protected.7z",
            {"secret.txt": b"secret"},
            password="test-password",
            header_encryption=True,
        )
        make_zip(
            marker_input / "bad.zip",
            {"password.txt": b"crc content"},
            zipfile.ZIP_STORED,
        )
        data = bytearray((marker_input / "bad.zip").read_bytes())
        central = data.index(b"PK\x01\x02")
        crc = struct.unpack_from("<I", data, central + 16)[0] ^ 1
        struct.pack_into("<I", data, central + 16, crc)
        struct.pack_into("<I", data, 14, crc)
        (marker_input / "bad.zip").write_bytes(data)

        sys.path.insert(0, str(ROOT))
        sys.path.insert(0, str(ROOT / "tests"))
        from test_folder_review import FolderReviewTests

        grouping_input = base / "grouping-input"
        grouping_input.mkdir()
        grouping_data = FolderReviewTests().grouping_payloads()
        for filename, (payload, _) in grouping_data.items():
            (grouping_input / filename).write_bytes(payload)
        make_zip(
            grouping_input / "nested.zip",
            {name: payload for name, (payload, _) in grouping_data.items()},
        )
        make_seven(
            grouping_input / "nested.7z",
            {name: payload for name, (payload, _) in grouping_data.items()},
        )

        inside_input = base / "inside-report-input"
        inside_input.mkdir()
        (inside_input / "data.txt").write_bytes(b"unchanged input")

        cases = [
            ("inside-report", inside_input, [], 0, output),
            ("logical-grouping", grouping_input, [], 0, output),
            ("presentation-markers", marker_input, [], 2, output),
            ("hashes-and-folders", hash_input, ["--full-metadata"], 0, output),
            ("complete", complete, [], 0, output),
            ("package-policy", leaf_input, [], 0, output),
            ("package-policy-full", leaf_input, ["--full-metadata"], 0, output),
            ("mixed", root, [], 2, output),
            ("full-metadata", root, ["--full-metadata"], 2, output),
            ("row-limit", root, ["--max-rows", "2"], 2, output),
            ("byte-limit", root, ["--max-expanded-gib", "0.000000001"], 2, output),
            ("time-limit", root, ["--max-archive-seconds", "0.000000001"], 2, output),
            ("codecs", codecs, [], 0, output),
            ("long-output", complete, [], 0, long_output),
            ("concurrency", root, ["--workers", "16"], 2, output),
        ] + rar_cases
        # Existing administrative share only: create no share or credentials.
        drive, tail = os.path.splitdrive(str(complete))
        unc = Path(f"\\\\localhost\\{drive[0]}$" + tail)
        if unc.is_dir():
            cases.append(("UNC", unc, [], 0, output))
            cases.append(("SMB-report", unc, [], 0, output))
        else:
            results["UNC"] = "not available to this logon session"
        # Map only our generated corpus, then remove the temporary mapping.
        mapped = next(
            (
                f"{letter}:"
                for letter in "ZYXWVUT"
                if not os.path.exists(f"{letter}:\\")
            ),
            None,
        )
        if mapped:
            subprocess.run(
                ["subst.exe", mapped, str(complete)], check=True, capture_output=True
            )
            cases.append(("mapped-drive", Path(mapped + "\\"), [], 0, output))
        try:
            input_folders = [
                root,
                complete,
                codecs,
                leaf_input,
                hash_input,
                marker_input,
                grouping_input,
                inside_input,
                *[case[1] for case in rar_cases],
            ]
            before = {str(folder): snapshot(folder) for folder in input_folders}
            for name, source, flags, expected, destination in cases:
                baseline = None
                for label, command, env in (
                    ("source", [sys.executable, str(ROOT / "folder_review.py")], None),
                    ("exe", [str(exe)], minimal_env),
                ):
                    target = (
                        (complete if name == "SMB-report" else source)
                        / "folder-review.xlsx"
                        if name in {"inside-report", "SMB-report"}
                        else destination / f"{name}-{label}.xlsx"
                    )
                    invocation = (
                        command
                        + [
                            "--root",
                            str(source),
                            "--output",
                            str(target),
                            "--workers",
                            "1",
                        ]
                        + flags
                    )
                    run = subprocess.run(
                        invocation,
                        capture_output=True,
                        text=True,
                        env=env,
                        timeout=120,
                        check=False,
                    )
                    assert run.returncode == expected, (
                        name,
                        label,
                        run.returncode,
                        run.stdout,
                        run.stderr,
                    )
                    current, status, count = normalized_workbook(target)
                    if name == "logical-grouping":
                        workbook = load_workbook(target, read_only=True, data_only=True)
                        try:
                            sheet = workbook["folder_review"]
                            headings = next(
                                sheet.iter_rows(min_row=6, max_row=6, values_only=True)
                            )
                            records = [
                                dict(zip(headings, row))
                                for row in sheet.iter_rows(min_row=7, values_only=True)
                            ]
                            assert len(records) == len(grouping_data) * 3
                            for row in records:
                                assert (
                                    row["Type group"],
                                    row["Type subgroup"],
                                ) == grouping_data[row["File name"]][1]
                                assert row["last"] == 1 and row["Virtual row"] == "No"
                        finally:
                            workbook.close()
                    if name == "presentation-markers":
                        workbook = load_workbook(target, read_only=True, data_only=True)
                        try:
                            sheet = workbook["folder_review"]
                            headings = next(
                                sheet.iter_rows(min_row=6, max_row=6, values_only=True)
                            )
                            records = [
                                dict(zip(headings, row))
                                for row in sheet.iter_rows(min_row=7, values_only=True)
                            ]
                            assert {row["File name"] for row in records} == {
                                "dummy.dummy",
                                "password.password",
                                "corrupt.corrupt",
                            }
                            assert all(
                                row["last"] == 1
                                and row["Virtual row"] == "Yes"
                                and row["Duplicate group"] is None
                                for row in records
                            )
                        finally:
                            workbook.close()
                    if name == "hashes-and-folders":
                        workbook = load_workbook(target, read_only=True, data_only=True)
                        try:
                            sheet = workbook["folder_review"]
                            headings = next(
                                sheet.iter_rows(min_row=6, max_row=6, values_only=True)
                            )
                            values = list(sheet.iter_rows(min_row=7, values_only=True))
                            records = [dict(zip(headings, row)) for row in values]
                            markers = [
                                row for row in records if row["Virtual row"] == "Yes"
                            ]
                            assert len(markers) == 2
                            assert all(
                                row["File name"] == "dummy.dummy"
                                and row["Content hash (BLAKE3)"] is None
                                and not row["Extension"]
                                for row in markers
                            )
                            assert all(row["last"] == 1 for row in records)
                            groups = [
                                row["Duplicate group"]
                                for row in records
                                if row["Duplicate group"] is not None
                            ]
                            assert sorted(set(groups)) == [1] and len(groups) == 4
                            assert (
                                max(groups.count(group) for group in set(groups)) == 4
                            )
                            for row in records:
                                leaf = row["File name"]
                                if leaf in hash_names | {"inside.txt"}:
                                    assert (
                                        row["Content hash (BLAKE3)"]
                                        == blake3(hash_payload).hexdigest()
                                    )
                                elif leaf == "empty-file":
                                    assert (
                                        row["Content hash (BLAKE3)"]
                                        == blake3(b"").hexdigest()
                                    )
                                if leaf == "1.2.3-cp313-cp313-win_amd64":
                                    assert not row["Extension"]
                                    assert row["Name without extension"] == leaf
                                    assert (
                                        row["Raw suffix"] == "3-cp313-cp313-win_amd64"
                                    )
                        finally:
                            workbook.close()
                    if name.startswith("package-policy"):
                        assert count == expected_leaf_rows, (
                            name,
                            label,
                            count,
                            expected_leaf_rows,
                        )
                        workbook = load_workbook(target, read_only=True, data_only=True)
                        try:
                            headers = next(
                                workbook["folder_review"].iter_rows(
                                    min_row=6, max_row=6, values_only=True
                                )
                            )
                            path_column = headers.index("Relative path")
                            filename_column = headers.index("File name")
                            for row in workbook["folder_review"].iter_rows(
                                min_row=7, values_only=True
                            ):
                                assert row[filename_column] in leaf_names | {
                                    "document.zip"
                                }, row
                                assert "::visio" not in row[path_column].lower(), row
                                assert "should-not-appear" not in row[path_column], row
                        finally:
                            workbook.close()
                    assert status["Status"] == (
                        "COMPLETE" if expected == 0 else "PARTIAL"
                    )
                    if baseline is None:
                        baseline = current
                    else:
                        differences = [
                            member
                            for member in current
                            if current[member] != baseline.get(member)
                        ]
                        assert current.keys() == baseline.keys() and not differences, (
                            name,
                            differences,
                        )
                        results["cases"].append(
                            {
                                "name": name,
                                "exit": expected,
                                "rows": count,
                                "all_workbook_members_match": True,
                            }
                        )
                if name in {"inside-report", "SMB-report"}:
                    target.unlink()
                assert not list(staging.glob("_MEI*")), (
                    "Frozen runtime directory leaked"
                )
                assert not list(staging.glob("folder-review-*")), (
                    "Nested archive workspace leaked"
                )
            after = {str(folder): snapshot(folder) for folder in input_folders}
            # Creating/removing the explicitly requested report changes its parent
            # directory timestamp; every existing file byte/metadata stays checked.
            for folder in (inside_input, complete):
                after[str(folder)]["."]["mtime_ns"] = before[str(folder)]["."][
                    "mtime_ns"
                ]
            assert before == after, "Source input bytes or unrelated metadata changed"

            # Copy only the EXE to a path with spaces and non-ASCII characters.
            import shutil

            portable = base / "portable folder-кириллица"
            portable.mkdir()
            relocated = portable / "folder-review.exe"
            shutil.copyfile(exe, relocated)
            run = subprocess.run(
                [str(relocated), "--root", str(complete)],
                cwd=output,
                env=minimal_env,
                capture_output=True,
                text=True,
                timeout=120,
                check=False,
            )
            assert run.returncode == 0, run.stdout + run.stderr
            assert (portable / "folder-review.xlsx").is_file()
            assert not (output / "folder-review.xlsx").exists()
            assert not list(staging.glob("_MEI*"))
            # The executable and its default report may be inside the input.
            for _ in range(2):
                run = subprocess.run(
                    [str(relocated), "--root", str(portable)],
                    cwd=output,
                    env=minimal_env,
                    capture_output=True,
                    text=True,
                    timeout=120,
                    check=False,
                )
                assert run.returncode == 0, run.stdout + run.stderr
                book = load_workbook(portable / "folder-review.xlsx", read_only=True)
                try:
                    sheet = book["folder_review"]
                    headings = next(
                        sheet.iter_rows(min_row=6, max_row=6, values_only=True)
                    )
                    names = [
                        row[headings.index("File name")]
                        for row in sheet.iter_rows(min_row=7, values_only=True)
                    ]
                    assert names == ["folder-review.exe"], names
                finally:
                    book.close()
            for flags, expected in (
                (["--root", str(base / "absent")], 1),
                (["--root", str(complete), "--output", str(output / "invalid.csv")], 1),
                (["--workers", "0"], 2),
                (["--help"], 0),
                (["--cli"], 1),
            ):
                old = subprocess.run(
                    [sys.executable, str(ROOT / "folder_review.py"), *flags],
                    capture_output=True,
                    text=True,
                    timeout=120,
                    check=False,
                )
                new = subprocess.run(
                    [str(exe), *flags],
                    capture_output=True,
                    text=True,
                    env=minimal_env,
                    timeout=120,
                    check=False,
                )
                assert old.returncode == new.returncode == expected, (
                    flags,
                    old.stderr,
                    new.stderr,
                )
            results["relocation_default_output_and_cli_errors"] = True
            results["runtime_cleanup"] = True
            results["source_input_bytes_unchanged"] = True
            results["source_tree_mtimes_attributes_and_links_unchanged"] = True
            results["sevenzip_filter_configurations"] = list(configurations)
        finally:
            subprocess.run(
                ["attrib.exe", "-H", "-R", str(root / "desktop.ini")], check=True
            )
            if junction.exists():
                os.rmdir(junction)
            if mapped:
                subprocess.run(
                    ["subst.exe", mapped, "/D"], check=True, capture_output=True
                )
    args.results.parent.mkdir(parents=True, exist_ok=True)
    args.results.write_text(json.dumps(results, indent=2), encoding="utf-8")
    print(json.dumps(results, indent=2))


if __name__ == "__main__":
    main()
