"""Notepad settings and the portable layout reuse the unchanged scanner."""

import contextlib
import io
import os
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

import folder_review as core


class SettingsTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.path = self.root / "folder-review.txt"
        self.location = mock.patch.object(
            core, "application_directory", return_value=self.root
        )
        self.location.start()
        self.addCleanup(self.location.stop)

    def test_literal_unicode_paths_encodings_and_all_options(self):
        target = r"\\server\share\пример 100% # data"
        text = (
            '[review]\nfolder = "' + target + '"\nduplicates = off\nworkers = 4\n'
            "report = result.xlsx\nfull_metadata = yes\ndiagnostics = no\n"
            "max_expanded_gib = 0\nmax_archive_seconds = 0\nmax_rows = 123\n"
        )
        for encoding in ("utf-8", "utf-8-sig", "utf-16"):
            with self.subTest(encoding=encoding):
                self.path.write_text(text, encoding=encoding)
                with mock.patch.object(core, "execute_review", return_value=2) as run:
                    self.assertEqual(core.run_settings(self.path), 2)
                args, selected, output, _ = run.call_args.args
                self.assertEqual(selected, target)
                self.assertEqual(output.parent, self.root)
                self.assertRegex(
                    output.name,
                    r"^result-\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2}(?:_\d{3})?\.xlsx$",
                )
                self.assertEqual(
                    (args.duplicates, args.workers, args.full_metadata, args.no_log),
                    ("off", 4, True, True),
                )
                self.assertEqual(
                    (args.max_expanded_gib, args.max_archive_seconds, args.max_rows),
                    (0, 0, 123),
                )

    def test_defaults_match_existing_scanner(self):
        self.path.write_text("[review]\nfolder = input\n")
        with mock.patch.object(core, "execute_review", return_value=0) as run:
            self.assertEqual(core.run_settings(self.path), 0)
        args = run.call_args.args[0]
        original = core.parse_args([])
        for name in (
            "duplicates",
            "workers",
            "full_metadata",
            "no_log",
            "max_expanded_gib",
            "max_archive_seconds",
            "max_rows",
        ):
            self.assertEqual(getattr(args, name), getattr(original, name), name)

    def test_invalid_settings_never_start_a_scan(self):
        variants = [
            "[wrong]\nfolder=x",
            "[DEFAULT]\nx=y\n[review]\nfolder=x",
            "[review]\nfolder=",
            "[review]\nfolder=x\nunknown=y",
            "[review]\nfolder=x\nfolder=y",
            "[review]\nfolder=x\nworkers=0",
            "[review]\nfolder=x\nduplicates=guess",
            "[review]\nfolder=x\nfull_metadata=perhaps",
            "[review]\nfolder=x\nmax_expanded_gib=nan",
            "[review]\nfolder=x\nmax_archive_seconds=inf",
            "[review]\nfolder=x\nmax_rows=0",
        ]
        for name in (
            "../escape.xlsx",
            r"C:\escape.xlsx",
            r"folder\escape.xlsx",
            "bad.csv",
            "",
        ):
            variants.append("[review]\nfolder=x\nreport=" + name)
        for text in variants:
            self.path.write_text(text)
            with (
                self.subTest(text=text),
                mock.patch.object(core, "execute_review") as run,
            ):
                with self.assertRaises((ValueError, core.argparse.ArgumentTypeError)):
                    core.run_settings(self.path)
                run.assert_not_called()

    def test_no_arguments_use_settings_without_a_gui(self):
        self.path.write_text("[review]\nfolder=input")
        with (
            mock.patch.object(core.sys, "argv", ["folder_review.py"]),
            mock.patch.object(core, "run_settings", return_value=2) as run,
        ):
            self.assertEqual(core.main(), 2)
        run.assert_called_once_with(self.path)

    def test_packaged_self_scan_skips_support_and_support_input_is_rejected(self):
        support = self.root / "support"
        support.mkdir()
        (support / "code.py").write_text("code")
        (self.root / "leaf.txt").write_bytes(b"unchanged")
        with (
            mock.patch.object(core, "_support_root", support),
            contextlib.redirect_stdout(io.StringIO()),
        ):
            with (
                self.assertRaisesRegex(ValueError, "support folder"),
                core.safe_runtime(str(support), self.root / "report.xlsx"),
            ):
                self.fail("Support input accepted")
            self.assertFalse((support / "state").exists())
            self.path.write_text("[review]\nfolder=" + str(self.root))
            self.assertEqual(core.run_settings(self.path), 0)
        self.assertEqual((self.root / "leaf.txt").read_bytes(), b"unchanged")
        self.assertFalse(list((support / "state").glob("work-*")))
        self.assertFalse((support / "state/logs").exists())
        book = load_workbook(
            next(self.root.glob("folder-review-*.xlsx")), read_only=True
        )
        try:
            rows = list(book["folder_review"].iter_rows(min_row=6, values_only=True))
            names = [r[rows[0].index("File name")] for r in rows[1:]]
            self.assertEqual(set(names), {"leaf.txt", "folder-review.txt"})
        finally:
            book.close()

    def test_only_the_fixed_local_backend_is_allowed_inside_self_scan(self):
        support = self.root / "support"
        backend = support / "state/backend/7z.exe"
        backend.parent.mkdir(parents=True)
        backend.write_bytes(b"fixture")
        with (
            mock.patch.object(core, "_support_root", support),
            core.safe_runtime(str(self.root), self.root / "report.xlsx"),
            mock.patch.object(core.shutil, "which", return_value=str(backend)),
        ):
            core.runtime_audit(
                "subprocess.Popen", (str(backend), [str(backend), "i"], None, None)
            )
            with self.assertRaises(PermissionError):
                core.runtime_audit(
                    "subprocess.Popen",
                    (str(backend), [str(backend), "e", "input.zip"], None, None),
                )

    def test_linked_log_directory_is_rejected_without_outside_writes(self):
        source = self.root / "input"
        source.mkdir()
        outside = self.root / "outside"
        outside.mkdir()
        link = core.portable_directory() / "logs"
        if os.name == "nt":
            subprocess.run(
                ["cmd.exe", "/c", "mklink", "/J", str(link), str(outside)],
                check=True,
                capture_output=True,
            )
        else:
            link.symlink_to(outside, target_is_directory=True)
        args = core.parse_args(["--log", str(link / "escape.jsonl")])
        try:
            with (
                core.safe_runtime(str(source), self.root / "report.xlsx"),
                self.assertRaisesRegex(ValueError, "normal directory"),
                core.performance_log(args, str(source)),
            ):
                self.fail("A linked log directory was accepted")
            self.assertEqual(list(outside.iterdir()), [])
        finally:
            if os.name == "nt":
                link.rmdir()
            else:
                link.unlink()


if __name__ == "__main__":
    unittest.main()
