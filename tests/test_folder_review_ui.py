"""Desktop acceptance: controls, background work and shared scanner behavior."""

from __future__ import annotations

import tempfile
import threading
import time
import tkinter as tk
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

import folder_review as core
from folder_review_ui import DUPLICATE_MODES, Application


class DesktopTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary.name)
        mock.patch.object(
            core, "application_directory", return_value=self.directory
        ).start()
        self.root = self.directory / "input folder-пример"
        self.root.mkdir()
        self.window = tk.Tk()
        self.window.withdraw()
        self.app = Application(self.window, self.directory)
        self.dialog = mock.patch("folder_review_ui.messagebox.showerror").start()

    def tearDown(self):
        self.wait()
        self.app.close()
        mock.patch.stopall()
        self.temporary.cleanup()

    def wait(self):
        deadline = time.monotonic() + 15
        while self.app.busy and time.monotonic() < deadline:
            self.window.update()
            time.sleep(0.01)
        self.window.update()
        self.assertFalse(self.app.busy)

    def test_blank_folder_disables_start_and_browse_accepts_selection(self):
        self.assertTrue(self.app.start_button.instate(["disabled"]))
        with mock.patch(
            "folder_review_ui.filedialog.askdirectory", return_value=str(self.root)
        ):
            self.app.browse_button.invoke()
        self.assertEqual(self.app.folder.get(), str(self.root))
        self.assertFalse(self.app.start_button.instate(["disabled"]))
        self.assertEqual(self.app.output.get(), str(self.directory / core.OUTPUT_FILE))
        self.assertEqual(str(self.app.output_entry.cget("state")), "readonly")
        self.assertFalse(hasattr(self.app, "save_button"))

    def test_complete_report_from_quoted_pasted_path_is_created_in_background(self):
        source = self.root / "file.txt"
        source.write_bytes(b"unchanged")
        self.app.folder.set('"' + str(self.root) + '"')
        self.app.start_button.invoke()
        self.assertTrue(self.app.busy)
        self.assertEqual(self.app.progress.winfo_manager(), "grid")
        self.assertTrue(self.app.browse_button.instate(["disabled"]))
        self.wait()
        self.assertIn("COMPLETE", self.app.status.get())
        self.assertEqual(self.app.progress.winfo_manager(), "")
        self.assertFalse(self.app.start_button.instate(["disabled"]))
        self.assertEqual(source.read_bytes(), b"unchanged")
        self.assertIn("Workers:", self.app.log.get("1.0", "end"))
        book = load_workbook(self.app.last_report, read_only=True)
        try:
            self.assertEqual(book["scan_status"]["B2"].value, "COMPLETE")
        finally:
            book.close()
        self.dialog.assert_not_called()

    def test_partial_report_is_explicit_and_not_reported_as_complete(self):
        (self.root / "bad.zip").write_bytes(b"PK\x03\x04bad")
        self.app.folder.set(str(self.root))
        self.app.start_button.invoke()
        self.wait()
        self.assertIn("PARTIAL", self.app.status.get())
        self.assertIsNotNone(self.app.last_report)
        self.dialog.assert_not_called()

    def test_all_duplicate_modes_reach_the_shared_scanner(self):
        (self.root / "first.txt").write_bytes(b"identical contents")
        (self.root / "second.txt").write_bytes(b"identical contents")
        self.app.folder.set(str(self.root))
        self.assertEqual(DUPLICATE_MODES[self.app.duplicate_mode.get()], "filtered")
        for label, mode in DUPLICATE_MODES.items():
            with self.subTest(mode=mode):
                self.app.duplicate_mode.set(label)
                self.app.start()
                self.assertEqual(str(self.app.duplicate_box.cget("state")), "disabled")
                self.wait()
                self.assertEqual(str(self.app.duplicate_box.cget("state")), "readonly")
                book = load_workbook(self.app.last_report, read_only=True)
                try:
                    status = {
                        row[0]: row[1]
                        for row in book["scan_status"].iter_rows(values_only=True)
                        if row[0] is not None
                    }
                    self.assertEqual(
                        status["Duplicate groups"], 0 if mode == "off" else 1
                    )
                    self.assertEqual(
                        status["Rows with fingerprints"], 0 if mode == "off" else 2
                    )
                finally:
                    book.close()

    def test_failure_restores_controls_and_preserves_last_report(self):
        self.app.folder.set(str(self.root))
        self.app.start()
        self.wait()
        previous = self.app.last_report
        old_bytes = previous.read_bytes()
        self.app.folder.set(str(self.directory / "absent"))
        self.app.start()
        self.wait()
        self.assertEqual(self.app.last_report, previous)
        self.assertEqual(previous.read_bytes(), old_bytes)
        self.assertIn("not saved", self.app.status.get())
        self.assertFalse(self.app.start_button.instate(["disabled"]))
        self.dialog.assert_called_once()

    def test_validation_is_on_worker_and_close_waits_for_it(self):
        entered, release = threading.Event(), threading.Event()
        caller = threading.get_ident()

        def delayed(*args, **kwargs):
            self.assertNotEqual(threading.get_ident(), caller)
            entered.set()
            release.wait(5)
            return 0

        self.app.folder.set(str(self.root))
        with mock.patch.object(core, "review_folder", side_effect=delayed):
            self.app.start()
            self.assertTrue(entered.wait(1))
            self.app.close()
            self.assertTrue(self.window.winfo_exists())
            self.assertIn("finish", self.app.status.get())
            release.set()
            self.wait()

    def test_thread_start_failure_does_not_leave_busy_controls(self):
        self.app.folder.set(str(self.root))
        with mock.patch(
            "folder_review_ui.threading.Thread.start",
            side_effect=RuntimeError("cannot start"),
        ):
            self.app.start()
            self.wait()
        self.assertIn("not saved", self.app.status.get())
        self.assertFalse(self.app.start_button.instate(["disabled"]))
        self.dialog.assert_called_once()

    def test_small_window_keeps_footer_visible_and_scrolls_card(self):
        self.window.minsize(760, 440)
        self.window.geometry("760x440")
        self.window.deiconify()
        self.window.update()
        self.assertEqual(self.app.scrollbar.winfo_manager(), "grid")
        self.assertLess(self.app.canvas.yview()[1], 1)
        self.assertLessEqual(
            self.app.open_button.winfo_rooty() + self.app.open_button.winfo_height(),
            self.window.winfo_rooty() + self.window.winfo_height(),
        )

    def test_no_arguments_select_gui_and_cli_arguments_preserve_cli(self):
        with (
            mock.patch.object(core.sys, "argv", ["folder_review.py"]),
            mock.patch("folder_review_ui.launch", return_value=0) as gui,
        ):
            self.assertEqual(core.main(), 0)
            gui.assert_called_once()
        with (
            mock.patch.object(core.sys, "argv", ["folder_review.py", "--cli"]),
            mock.patch.object(core, "run", return_value=2) as cli,
        ):
            self.assertEqual(core.main(), 2)
            cli.assert_called_once()


if __name__ == "__main__":
    unittest.main()
