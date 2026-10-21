"""Speed-first defaults, bounded approximate duplicate reads and saved reports."""

import contextlib
import io
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import py7zr
from openpyxl import load_workbook

import folder_change_review as changes
import folder_review as core
from duplicate_sampling import SAMPLE_BYTES, SampleCapture
from review_state import load_snapshot


class FastModeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.tool = self.base / "tool"
        self.tool.mkdir()
        self.root = self.base / "input"
        self.root.mkdir()
        location = patch.object(core, "application_directory", return_value=self.tool)
        location.start()
        self.addCleanup(location.stop)

    def inspect(self, mode):
        files, _ = core.scan_tree(str(self.root), set(), 2)
        with contextlib.redirect_stdout(io.StringIO()):
            return core.inspect_files(files, 2, duplicates=mode)

    def test_normal_is_bounded_and_never_claims_exact_equality(self):
        data = bytearray(b"a" * (SAMPLE_BYTES * 16))
        (self.root / "a.bin").write_bytes(data)
        data[SAMPLE_BYTES * 2] = ord("Z")  # outside all three samples
        (self.root / "b.bin").write_bytes(data)
        with patch.object(
            core, "file_digest", side_effect=AssertionError("Full read in Normal")
        ):
            normal = self.inspect("normal")
        self.assertEqual(normal[0].sample_hash, normal[1].sample_hash)
        self.assertTrue(all(row.content_hash is None for row in normal))
        full = self.inspect("full")
        self.assertNotEqual(full[0].content_hash, full[1].content_hash)
        report = self.tool / "normal.xlsx"
        with core.safe_runtime(str(self.root), report):
            core.write_workbook(normal, report, str(self.root))
        book = load_workbook(report)
        self.addCleanup(book.close)
        headers = [c.value for c in book["folder_review"][6]]
        self.assertIn("Possible duplicate group", headers)
        self.assertNotIn("Duplicate group", headers)
        group = headers.index("Possible duplicate group")
        self.assertEqual(
            [
                r[group]
                for r in book["folder_review"].iter_rows(min_row=7, values_only=True)
            ],
            [1, 1],
        )
        self.assertIn(
            "NOT confirmed",
            dict(
                book["scan_status"].iter_rows(
                    min_row=2, max_row=30, max_col=2, values_only=True
                )
            )["Content fingerprints"],
        )

    def test_normal_file_reads_at_most_24_kib_and_checks_mutations(self):
        file = self.root / "a.bin"
        file.write_bytes(b"a" * (1024 * 1024))
        actual_open = io.open
        requests = []

        class Reader:
            def __init__(self, *args, **kwargs):
                self.stream = actual_open(*args, **kwargs)

            def __enter__(self):
                return self

            def __exit__(self, *args):
                self.stream.close()

            def seek(self, *args):
                return self.stream.seek(*args)

            def read(self, count):
                value = self.stream.read(min(count, 127))
                requests.append(len(value))
                return value

        with patch.object(core, "open", side_effect=Reader, create=True):
            digest = core.file_sample_digest(
                str(file), file.stat().st_size, file.stat().st_mtime_ns
            )
        self.assertEqual(len(digest), 32)
        self.assertEqual(sum(requests), 3 * SAMPLE_BYTES)
        with self.assertRaises(OSError):
            core.file_sample_digest(
                str(file), file.stat().st_size, file.stat().st_mtime_ns - 1
            )

    def test_stream_samples_equal_file_samples_for_short_and_large_members(self):
        for size in (
            0,
            1,
            SAMPLE_BYTES - 1,
            SAMPLE_BYTES + 1,
            3 * SAMPLE_BYTES,
            16 * SAMPLE_BYTES,
        ):
            data = (bytes(range(256)) * ((size + 255) // 256))[:size]
            file = self.root / "a.bin"
            file.write_bytes(data)
            capture = SampleCapture(size)
            for offset in range(0, size, 137):
                capture.update(data[offset : offset + 137])
            self.assertEqual(capture.digest(), core.file_sample_digest(str(file), size))
            self.assertLessEqual(sum(capture.counts), 3 * SAMPLE_BYTES)
        bad = SampleCapture(20)
        bad.update(b"short")
        with self.assertRaises(OSError):
            bad.digest()

    def test_archive_members_and_physical_files_share_possible_groups(self):
        payload = b"same payload" * 20000
        (self.root / "physical.bin").write_bytes(payload)
        with zipfile.ZipFile(self.root / "zip.zip", "w") as z:
            z.writestr("member.bin", payload)
        with py7zr.SevenZipFile(self.root / "seven.7z", "w") as z:
            z.writestr(payload, "member.bin")
        with patch.object(
            core, "file_digest", side_effect=AssertionError("Full hash in Normal")
        ):
            rows = self.inspect("normal")
        self.assertEqual(len(rows), 3)
        self.assertEqual(len({row.sample_hash for row in rows}), 1)
        self.assertTrue(
            all(row.sample_hash is not None and row.error is None for row in rows)
        )
        self.assertTrue(all(row.content_hash is None for row in rows))
        self.assertEqual(
            {row.integrity_status for row in rows if row.in_archive == "Yes"},
            {core.INTEGRITY_VERIFIED},
        )

    def test_quick_and_off_skip_all_duplicate_reads(self):
        for name in ("a.bin", "b.bin"):
            (self.root / name).write_bytes(b"same")
        with (
            patch.object(core, "file_digest", side_effect=AssertionError("hash")),
            patch.object(
                core, "file_sample_digest", side_effect=AssertionError("sample")
            ),
        ):
            for mode in ("quick", "off"):
                self.assertTrue(
                    all(
                        row.content_hash is None and row.sample_hash is None
                        for row in self.inspect(mode)
                    )
                )

    def test_default_settings_create_no_logs_and_preserve_all_timestamped_states(self):
        self.root = self.tool
        (self.root / "business.txt").write_bytes(b"unchanged")
        config = self.tool / "folder-review.txt"
        config.write_text("[review]\nfolder=" + str(self.root))
        stamp = datetime(2026, 10, 21, 14, 5, 9)
        with (
            patch.object(core, "local_now", return_value=stamp),
            contextlib.redirect_stdout(io.StringIO()),
        ):
            for _ in range(3):
                self.assertEqual(core.run_settings(config), 0)
        files = sorted(self.tool.glob("folder-review-*.xlsx"))
        self.assertEqual(len(files), 3)
        self.assertEqual(
            {f.name for f in files},
            {
                "folder-review-2026-10-21_14-05-09.xlsx",
                "folder-review-2026-10-21_14-05-09_001.xlsx",
                "folder-review-2026-10-21_14-05-09_002.xlsx",
            },
        )
        for path in files:
            self.assertEqual(
                [r["parts"] for r in load_snapshot(path)["files"]], [["business.txt"]]
            )
        compare = self.tool / "folder-change-review.txt"
        compare.write_text(f"[compare]\nbefore={files[0]}\nafter={files[1]}")
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(changes.run_settings(compare), 0)
        self.assertFalse((self.tool / "folder-review-data/logs").exists())
        self.assertTrue(core.parse_args([]).no_log)
        self.assertEqual(core.parse_args([]).duplicates, "quick")

    def test_report_history_exclusion_scales_beyond_old_scope_limit(self):
        self.root = self.tool
        (self.root / "business.txt").write_bytes(b"good")
        for n in range(80):
            (self.root / f"folder-review-2026-10-21_14-05-09_{n:03d}.xlsx").write_bytes(
                b"report artifact"
            )
        config = self.tool / "folder-review.txt"
        config.write_text("[review]\nfolder=" + str(self.root))
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(core.run_settings(config), 0)
        output = max(
            self.tool.glob("folder-review-*.xlsx"), key=lambda f: f.stat().st_mtime_ns
        )
        self.assertEqual(load_snapshot(output)["file_count"], 1)
        # The filename pattern is reserved only beside this application.
        (self.base / "input/folder-review-2026-10-21_14-05-09.xlsx").write_bytes(
            b"ordinary file"
        )
        files, _ = core.scan_tree(
            str(self.base / "input"),
            core.scan_exclusions(output, str(self.base / "input")),
            1,
        )
        self.assertEqual(len(files), 1)
