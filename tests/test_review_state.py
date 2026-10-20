"""End-to-end inventory states, comparison coverage and portable input safety."""

import contextlib
import copy
import hashlib
import io
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import folder_change_review as changes
import folder_review as core
import review_state as state


class WorkbookStateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.tool = self.base / "tool"
        self.tool.mkdir()
        self.root = self.base / "project"
        self.root.mkdir()
        location = patch.object(core, "application_directory", return_value=self.tool)
        location.start()
        self.addCleanup(location.stop)
        self.sequence = 0
        self.output = self.tool / "folder-change-review.xlsx"

    def capture(self, mode="quick", **options):
        self.sequence += 1
        output = self.tool / f"state-{self.sequence}.xlsx"
        args = core.parse_args(
            ["--duplicates", "off", "--no-log", "--state-mode", mode]
        )
        for name, value in options.items():
            setattr(args, name, value)
        with contextlib.redirect_stdout(io.StringIO()):
            code = core.execute_review(args, str(self.root), output)
        return state.load_snapshot(output), output, code

    def save(self, snapshot, name="custom.xlsx"):
        output = self.tool / name
        with (
            core.safe_runtime(str(self.root), output),
            core._report_workbook(output) as wb,
        ):
            wb.add_worksheet("folder_review")
            state.write_state(wb, snapshot)
        return output

    def compare(self, a, b):
        with contextlib.redirect_stdout(io.StringIO()):
            return changes.compare_files(a, b, self.output, diagnostics=False)

    def test_physical_state_exclusions_and_hidden_inventory_leaf_distinction(self):
        for name in ("~lock.docx", "Thumbs.db", "a.TMP", "a.TeMp"):
            (self.root / name).write_bytes(b"ignored")
        (self.root / "empty").mkdir()
        (self.root / "normal.txt").write_bytes(b"normal")
        with zipfile.ZipFile(self.root / "package.zip", "w") as z:
            z.writestr("inside.txt", "contents")
        with (
            patch.object(
                changes, "hash_file", side_effect=AssertionError("extra full read")
            ),
            patch.object(
                changes, "sample_file", side_effect=AssertionError("extra sample")
            ),
        ):
            snapshot, path, code = self.capture()
        self.assertEqual(code, 0)
        self.assertEqual(
            [r["parts"] for r in snapshot["files"]], [["normal.txt"], ["package.zip"]]
        )
        self.assertEqual(snapshot["statistics"]["content_bytes_read"], 0)
        book = load_workbook(path)
        self.addCleanup(book.close)
        self.assertEqual(book[state.SHEET].sheet_state, "hidden")
        names = [
            c.value for row in book["folder_review"].iter_rows(min_row=7) for c in row
        ]
        self.assertIn("inside.txt", names)
        self.assertIn("dummy.dummy", names)
        self.assertEqual(book["folder_review"].freeze_panes, "A7")

    def test_precise_timestamps_are_retained_without_excel_number_rounding(self):
        path = self.root / "a.txt"
        path.write_bytes(b"a")
        stamp = 1_700_000_000_123_456_700
        os.utime(path, ns=(stamp, stamp))
        a, pa, _ = self.capture()
        os.utime(path, ns=(stamp + 100, stamp + 100))
        b, pb, _ = self.capture()
        self.assertEqual(a["files"][0]["mtime_ns"], stamp)
        self.assertEqual(b["files"][0]["mtime_ns"], stamp + 100)
        self.assertEqual(self.compare(pa, pb), {"Changed": 1})

    def test_modes_preserve_original_detection_rules(self):
        file = self.root / "file.bin"
        file.write_bytes(b"a" * (changes.SAMPLE_SIZE * 12))
        stamp = file.stat().st_mtime_ns
        before = {mode: self.capture(mode)[1] for mode in changes.REVIEW_MODES}
        with file.open("r+b") as stream:
            stream.seek(changes.SAMPLE_SIZE * 2)
            stream.write(b"Z")
        os.utime(file, ns=(stamp, stamp))
        for mode in changes.REVIEW_MODES:
            snapshot, path, _ = self.capture(mode)
            self.assertEqual(
                self.compare(before[mode], path),
                {"Changed" if mode == "full" else "Unchanged": 1},
            )
            if mode == "full":
                self.assertEqual(
                    snapshot["files"][0]["sha256"],
                    hashlib.sha256(file.read_bytes()).hexdigest(),
                )
            if mode == "sampled":
                self.assertEqual(
                    snapshot["statistics"]["content_bytes_read"],
                    3 * changes.SAMPLE_SIZE,
                )

    def test_sample_modes_and_short_files_reuse_original_fingerprints(self):
        file = self.root / "a.bin"
        for size in (0, 1, changes.SAMPLE_SIZE, changes.SAMPLE_SIZE + 1):
            file.write_bytes(b"a" * size)
            snapshot, _, _ = self.capture("sampled")
            expected = changes.sample_file(str(file), file.stat())
            self.assertEqual(snapshot["files"][0]["sha256"], expected[0])
            self.assertEqual(snapshot["statistics"]["content_bytes_read"], expected[2])

    def test_hash_failure_creates_unknown_and_partial_without_losing_inventory_row(
        self,
    ):
        file = self.root / "a.txt"
        file.write_bytes(b"content")
        _, before, _ = self.capture("full")
        with patch.object(
            changes, "hash_file", side_effect=PermissionError(13, "Denied")
        ):
            snapshot, after, code = self.capture("full")
        self.assertEqual(code, 2)
        self.assertFalse(snapshot["complete"])
        self.assertEqual(snapshot["read_errors"][0]["kind"], "file")
        counts = self.compare(before, after)
        self.assertEqual(counts, {"Unknown": 1})
        self.assertTrue(counts.partial)
        book = load_workbook(self.output)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["Changes", "Read issues"])
        self.assertIn("Incomplete coverage", book["Changes"]["B3"].value)

    def test_directory_failure_marks_missing_descendants_unknown_but_not_siblings(self):
        for folder in ("Area", "Area2"):
            (self.root / folder).mkdir()
            (self.root / folder / "a.txt").write_bytes(b"a")
        _, before, _ = self.capture()
        original = core.os.scandir

        def denied(path):
            if not isinstance(path, int) and core.absolute_path(path) == str(
                self.root / "Area"
            ):
                raise PermissionError(13, "Denied")
            return original(path)

        with patch.object(core.os, "scandir", side_effect=denied):
            snapshot, after, code = self.capture()
        self.assertEqual(code, 2)
        self.assertEqual(self.compare(before, after), {"Unknown": 1, "Unchanged": 1})
        self.assertEqual(snapshot["read_errors"][0]["parts"], ["Area"])

    def test_root_failure_keeps_coverage_warning_even_with_no_known_filenames(self):
        original = core.os.scandir

        def denied(path):
            if not isinstance(path, int) and core.absolute_path(path) == str(self.root):
                raise PermissionError(13, "Denied")
            return original(path)

        with patch.object(core.os, "scandir", side_effect=denied):
            snapshot, path, code = self.capture()
        self.assertEqual(code, 2)
        self.assertEqual(snapshot["read_errors"][0]["parts"], [])
        counts = self.compare(path, path)
        self.assertTrue(counts.partial)
        self.assertEqual(counts, {})
        settings = self.tool / "folder-change-review.ini"
        settings.write_text(f"[compare]\nbefore={path}\nafter={path}\ndiagnostics=no")
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(changes.run_settings(settings), 2)

    def test_archive_errors_and_visible_row_limits_do_not_invent_physical_absence(self):
        for name in ("a.txt", "b.txt", "c.txt"):
            (self.root / name).write_bytes(b"a")
        (self.root / "broken.zip").write_bytes(b"PK\x03\x04broken")
        snapshot, _, code = self.capture(max_rows=1)
        self.assertEqual(code, 2)
        self.assertTrue(snapshot["complete"])
        self.assertEqual(snapshot["file_count"], 4)
        self.assertIn(["broken.zip"], [r["parts"] for r in snapshot["files"]])

    def test_comparison_never_reads_live_project_or_changes_input_states(self):
        (self.root / "a.txt").write_bytes(b"a")
        snapshot, _, _ = self.capture()
        snapshot["root"] = r"\\unavailable-server\share\project"
        snapshot["files"][0]["path"] = snapshot["root"] + r"\a.txt"
        path = self.save(snapshot)
        original = path.read_bytes()
        stamp = path.stat().st_mtime_ns
        real_stat = core.os.stat

        def guarded(candidate, *args, **kwargs):
            if isinstance(candidate, (str, os.PathLike)) and str(candidate).startswith(
                snapshot["root"]
            ):
                raise AssertionError("Live root contacted")
            return real_stat(candidate, *args, **kwargs)

        with patch.object(core.os, "stat", side_effect=guarded):
            self.assertEqual(self.compare(path, path), {"Unchanged": 1})
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual(path.stat().st_mtime_ns, stamp)

    def test_mode_mismatch_legacy_and_corrupt_states_preserve_previous_report(self):
        (self.root / "a.txt").write_bytes(b"a")
        _, quick, _ = self.capture()
        _, full, _ = self.capture("full")
        self.output.write_bytes(b"previous")
        with self.assertRaisesRegex(ValueError, "different review modes"):
            self.compare(quick, full)
        legacy = self.tool / "old.xlsx"
        with (
            core.safe_runtime(str(self.root), legacy),
            core._report_workbook(legacy) as wb,
        ):
            wb.add_worksheet("folder_review")
        with self.assertRaisesRegex(ValueError, "no comparison state"):
            self.compare(legacy, quick)
        corrupt = self.tool / "corrupt.xlsx"
        corrupt.write_bytes(b"corrupt")
        with self.assertRaises(ValueError):
            self.compare(corrupt, quick)
        self.assertEqual(self.output.read_bytes(), b"previous")

    def test_output_cannot_replace_input_or_escape_portable_folder(self):
        (self.root / "a.txt").write_bytes(b"a")
        _, path, _ = self.capture()
        original = path.read_bytes()
        for output in (path, self.root / "bad.xlsx"):
            with self.assertRaises(ValueError):
                changes.compare_files(path, path, output, diagnostics=False)
        self.assertEqual(path.read_bytes(), original)
        if os.name != "nt":
            alias = self.tool / "alias.xlsx"
            os.link(path, alias)
            with self.assertRaisesRegex(ValueError, "replace an input"):
                changes.compare_files(path, path, alias, diagnostics=False)

    def test_state_validation_rejects_traversal_duplicates_modes_and_invalid_coverage(
        self,
    ):
        (self.root / "a.txt").write_bytes(b"a")
        snapshot, _, _ = self.capture()
        variants = []
        for key, value in (
            ("version", 99),
            ("policy", "other"),
            ("complete", False),
            ("review_mode", "other"),
            ("excluded_tool_paths", ["../escape"]),
        ):
            bad = copy.deepcopy(snapshot)
            bad[key] = value
            variants.append(bad)
        bad = copy.deepcopy(snapshot)
        bad["files"][0]["parts"] = ["..", "escape"]
        variants.append(bad)
        bad = copy.deepcopy(snapshot)
        bad["files"] *= 2
        bad["file_count"] = 2
        variants.append(bad)
        bad = copy.deepcopy(snapshot)
        bad["files"][0]["mtime_ns"] = "rounded"
        variants.append(bad)
        for bad in variants:
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                state.validate_snapshot(bad)

    def test_state_write_limit_and_report_failure_are_atomic(self):
        (self.root / "a.txt").write_bytes(b"a")
        snapshot, _, _ = self.capture()
        output = self.tool / "custom.xlsx"
        output.write_bytes(b"previous")
        with patch.object(state, "MAX_JSON", 1), self.assertRaises(ValueError):
            self.save(snapshot)
        self.assertEqual(output.read_bytes(), b"previous")
        self.assertFalse(list(self.tool.glob(".folder-review-*.tmp.xlsx")))
        self.assertFalse(list((self.tool / "folder-review-data").glob("work-*")))

    def test_notepad_comparison_config_is_literal_and_separate(self):
        (self.root / "пример 100%.txt").write_bytes(b"a")
        _, path, _ = self.capture()
        config = self.tool / "folder-change-review.ini"
        for encoding in ("utf-8", "utf-8-sig", "utf-16"):
            config.write_text(
                f'[compare]\nbefore="{path}"\nafter="{path}"\nreport=changes.xlsx\ndiagnostics=no',
                encoding=encoding,
            )
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(changes.run_settings(config), 0)
        for text in (
            "[compare]\nbefore=relative.xlsx\nafter=relative.xlsx",
            "[compare]\nunknown=x",
            "[DEFAULT]\nx=y\n[compare]",
            "[compare]\nbefore=x\nbefore=y",
        ):
            config.write_text(text)
            with self.assertRaises(ValueError):
                changes.run_settings(config)

    def test_neutral_comparison_log_contains_only_aggregate_fields(self):
        (self.root / "private-file.txt").write_bytes(b"a")
        _, path, _ = self.capture()
        with contextlib.redirect_stdout(io.StringIO()):
            changes.compare_files(path, path, self.output)
        for log in (self.tool / "folder-review-data/logs").glob("*.jsonl"):
            text = log.read_text()
            self.assertNotIn("private-file", text)
            self.assertNotIn(str(self.root), text)
            for line in text.splitlines():
                json.loads(line)

    def test_source_read_only_and_network_guards_apply_to_comparison(self):
        import socket

        (self.root / "a.txt").write_bytes(b"a")
        _, path, _ = self.capture()
        original = changes.write_report

        def guarded(*args, **kwargs):
            for operation in (
                lambda: path.write_bytes(b"bad"),
                lambda: socket.socket(),
                lambda: (self.root / "bad").write_bytes(b"bad"),
            ):
                with self.assertRaises(PermissionError):
                    operation()
            return original(*args, **kwargs)

        with patch.object(changes, "write_report", side_effect=guarded):
            self.compare(path, path)

    def test_corrupt_embedded_text_lengths_and_entities_fail_without_replacing_report(
        self,
    ):
        (self.root / "a.txt").write_bytes(b"a")
        _, original, _ = self.capture()
        with zipfile.ZipFile(original) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        hidden = members["xl/worksheets/sheet4.xml"]
        variants = [
            hidden.replace(b"folder-review-state-v1", b"folder-review-state-v9"),
            hidden.replace(b"<is><t>", b"<f>1+1</f><is><t>", 1),
            hidden.replace(
                b"<worksheet ",
                b'<!DOCTYPE worksheet [<!ENTITY bad "bad">]><worksheet ',
                1,
            ),
            hidden.replace(b"</worksheet>", b"<row></worksheet>"),
        ]
        self.output.write_bytes(b"previous")
        for index, xml in enumerate(variants):
            bad = self.tool / f"tampered-{index}.xlsx"
            with zipfile.ZipFile(bad, "w") as archive:
                for name, data in members.items():
                    archive.writestr(
                        name, xml if name == "xl/worksheets/sheet4.xml" else data
                    )
            with self.subTest(index=index), self.assertRaises(ValueError):
                self.compare(bad, original)
            self.assertEqual(self.output.read_bytes(), b"previous")
