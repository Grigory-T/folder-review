"""Comparison/layout regressions migrated from folder-change-review (MIT)."""

import copy
import hashlib
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import folder_change_review as app
import folder_review as core


class MigratedComparisonTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name)
        self.root = self.work / "project"
        self.root.mkdir()
        location = patch.object(core, "application_directory", return_value=self.work)
        location.start()
        self.addCleanup(location.stop)
        self.complete = {"root": str(self.root), "files": []}

    def snapshot(self, name):
        files = []
        for path in sorted(self.root.rglob("*")):
            if path.is_file():
                files.append(
                    {
                        "path": str(path),
                        "parts": list(path.relative_to(self.root).parts),
                        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                        "size": path.stat().st_size,
                        "mtime_ns": path.stat().st_mtime_ns,
                    }
                )
        return {"root": str(self.root), "review_mode": "full", "files": files}

    def write_report(self, rows, output, **kwargs):
        with core.safe_runtime(str(self.root), output):
            app.write_report(rows, output, **kwargs)

    def test_union_hash_changes_moves_and_full_paths(self):
        for directory in ("one", "two"):
            (self.root / directory).mkdir()
            (self.root / directory / "same.pdf").write_bytes(b"same")
        (self.root / "changed.docx").write_bytes(b"before")
        (self.root / "deleted.pdf").write_bytes(b"delete")
        (self.root / "move.pdf").write_bytes(b"move")
        before = self.snapshot("A.fcrstate")
        (self.root / "changed.docx").write_bytes(b"AFTER!")  # same size, different hash
        (self.root / "deleted.pdf").unlink()
        (self.root / "added.pdf").write_bytes(b"add")
        (self.root / "move.pdf").rename(self.root / "one/move.pdf")
        os.utime(self.root / "one/same.pdf", (1_700_000_000, 1_700_000_000))
        after = self.snapshot("B.fcrstate")
        rows = app.compare_snapshots(before, after)
        self.assertEqual(len(rows), 7)
        self.assertEqual(
            {
                status: sum(row["Status"] == status for row in rows)
                for status in ("Added", "Deleted", "Changed", "Unchanged")
            },
            {"Added": 2, "Deleted": 2, "Changed": 1, "Unchanged": 2},
        )
        for row in rows:
            self.assertEqual(
                sum(row[name] or 0 for name in ("Added", "Deleted", "Changed")),
                int(row["Status"] != "Unchanged"),
            )
        self.assertIsNone(
            next(row for row in rows if row["FileName"] == "added.pdf")["SizeA"]
        )

    def test_excel_plain_numeric_flags_and_literal_names(self):
        (self.root / "=1+1.txt").write_bytes(b"text")
        (self.root / "00123.txt").write_bytes(b"text")
        (self.root / "пример-😀.pdf").write_bytes(b"pdf")
        after = self.snapshot("B.fcrstate")
        rows = app.compare_snapshots({"files": []}, after)
        output = self.work / "changes.xlsx"
        self.write_report(rows, output)
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["Changes"])
        sheet = book["Changes"]
        self.assertEqual(len(sheet.tables), 0)
        headers = [cell.value for cell in sheet[6]][1:]
        self.assertNotIn("Leaf", headers)
        self.assertNotIn("FullPath", headers)
        self.assertFalse(any("SHA" in name.upper() for name in headers))
        self.assertEqual(headers[0], "№пп")
        self.assertEqual(
            headers,
            [
                "№пп",
                "Level 1",
                "Leaf depth",
                "Name without extension",
                "File name",
                "Extension",
                "Size A (bytes)",
                "Size A (MiB)",
                "Size B (bytes)",
                "Size B (MiB)",
                "Size delta (bytes)",
                "Size delta (MiB)",
                "Status",
                "Added",
                "Deleted",
                "Changed",
                "last",
            ],
        )
        self.assertEqual(headers[-1], "last")
        self.assertEqual(sheet.max_row, 9)
        self.assertEqual(sheet.freeze_panes, "A7")
        self.assertEqual(sheet.sheet_view.zoomScale, 70)
        self.assertEqual(sheet.row_dimensions[6].height, 72)
        self.assertEqual(sheet["B6"].fill.fgColor.rgb, "FF000000")
        self.assertEqual(sheet["C6"].fill.fgColor.rgb, "FF00B0F0")
        self.assertEqual(
            list(
                next(sheet.iter_rows(min_row=4, max_row=4, min_col=2, values_only=True))
            ),
            list(range(1, len(headers) + 1)),
        )
        for cells in sheet.iter_rows(min_row=7, min_col=2):
            self.assertEqual(cells[headers.index("File name")].data_type, "s")
            self.assertEqual(cells[headers.index("Added")].value, 1)
            self.assertEqual(cells[headers.index("Added")].data_type, "n")
            self.assertIsNone(cells[headers.index("Deleted")].value)
            self.assertIsNone(cells[headers.index("Size A (MiB)")].value)
            self.assertEqual(
                cells[headers.index("Size B (MiB)")].number_format, "#,##0.000"
            )
            self.assertIn(cells[headers.index("Extension")].value, ("txt", "pdf"))
            self.assertEqual(cells[-1].value, 1)
            self.assertTrue(all(cell.fill.patternType is None for cell in cells))
            self.assertTrue(all(cell.border.left.style == "thin" for cell in cells))
            self.assertTrue(
                all(cell.data_type != "f" and cell.hyperlink is None for cell in cells)
            )

    def test_english_order_keeps_sizes_flags_hierarchy_and_names_aligned(self):
        (self.root / "nested").mkdir()
        (self.root / "nested/данные.tar.gz").write_bytes(b"before")
        (self.root / "deleted.pdf").write_bytes(b"deleted")
        (self.root / ".gitignore").write_bytes(b"")
        before = self.snapshot("A.fcrstate")
        (self.root / "nested/данные.tar.gz").write_bytes(b"x" * 1_234_567)
        (self.root / "deleted.pdf").unlink()
        (self.root / "zero").write_bytes(b"")
        after = self.snapshot("B.fcrstate")
        output = self.work / "aligned.xlsx"
        self.write_report(app.compare_snapshots(before, after), output)
        book = load_workbook(output)
        self.addCleanup(book.close)
        sheet = book["Changes"]
        self.assertEqual(
            sheet["B2"].value, f"Table: File changes in folder {self.root}"
        )
        headers = [cell.value for cell in sheet[6]][1:]
        self.assertEqual(
            headers,
            [
                "№пп",
                "Level 1",
                "Level 2",
                "Leaf depth",
                "Name without extension",
                "File name",
                "Extension",
                "Size A (bytes)",
                "Size A (MiB)",
                "Size B (bytes)",
                "Size B (MiB)",
                "Size delta (bytes)",
                "Size delta (MiB)",
                "Status",
                "Added",
                "Deleted",
                "Changed",
                "last",
            ],
        )
        records = {
            row[headers.index("File name")]: dict(zip(headers, row))
            for row in sheet.iter_rows(min_row=7, min_col=2, values_only=True)
        }
        changed = records["данные.tar.gz"]
        self.assertEqual(
            (changed["Level 1"], changed["Level 2"], changed["Leaf depth"]),
            ("nested", "данные.tar.gz", 2),
        )
        self.assertEqual(
            (changed["Name without extension"], changed["Extension"]),
            ("данные.tar", "gz"),
        )
        expected = {
            "данные.tar.gz": (6, 1_234_567, "Changed", None, None, 1),
            "deleted.pdf": (7, None, "Deleted", None, 1, None),
            "zero": (None, 0, "Added", 1, None, None),
            ".gitignore": (0, 0, "Unchanged", None, None, None),
        }
        for name, (size_a, size_b, status, added, deleted, changed) in expected.items():
            record = records[name]
            for label, size in (("Size A (MiB)", size_a), ("Size B (MiB)", size_b)):
                if size is None:
                    self.assertIsNone(record[label])
                else:
                    self.assertAlmostEqual(record[label], size / 1_048_576, delta=1e-14)
            self.assertEqual(
                (
                    record["Status"],
                    record["Added"],
                    record["Deleted"],
                    record["Changed"],
                ),
                (status, added, deleted, changed),
            )
            self.assertEqual(record["last"], 1)
        self.assertEqual(records[".gitignore"]["Name without extension"], ".gitignore")
        self.assertEqual(records[".gitignore"]["Extension"], "")

    def test_lca_uses_union_including_deleted_branch_and_single_file(self):
        paths = [
            str(self.root / "shared/Finance/a.pdf"),
            str(self.root / "shared/Financial/b.pdf"),
        ]
        rows = [{"FullPath": path, "parts": ["unused"]} for path in paths]
        ancestor, parts = app.report_hierarchy(rows)
        self.assertEqual(ancestor, str(self.root / "shared"))
        self.assertEqual(parts, [["Finance", "a.pdf"], ["Financial", "b.pdf"]])
        self.assertEqual(
            app.report_hierarchy(rows[:1]),
            (str(self.root / "shared/Finance"), [["a.pdf"]]),
        )
        windows = [
            {"FullPath": r"C:\Project\Common\one\a.pdf", "parts": []},
            {"FullPath": r"C:\Project\Common\two\b.pdf", "parts": []},
        ]
        self.assertEqual(
            app.report_hierarchy(windows),
            (r"C:\Project\Common", [["one", "a.pdf"], ["two", "b.pdf"]]),
        )

    def test_row_limit_accounts_for_six_header_rows(self):
        self.assertEqual(app.MAX_DATA_ROWS, 1_048_570)

    def test_windows_unc_scopes_match_ancestors_without_matching_siblings(self):
        root = r"\\server\share\Project"
        before = {
            "root": root,
            "files": [
                {
                    "path": root + r"\Area\Nested\one.pdf",
                    "parts": ["Area", "Nested", "one.pdf"],
                    "size": 1,
                    "sha256": "a" * 64,
                },
                {
                    "path": root + r"\Area2\two.pdf",
                    "parts": ["Area2", "two.pdf"],
                    "size": 1,
                    "sha256": "a" * 64,
                },
            ],
        }
        after = {
            "root": root,
            "files": [],
            "read_errors": [
                {"path": root + r"\Area", "parts": ["Area"], "kind": "directory"}
            ],
        }
        rows = app.compare_snapshots(before, after)
        self.assertEqual(
            {row["FileName"]: row["Status"] for row in rows},
            {"one.pdf": "Unknown", "two.pdf": "Deleted"},
        )

    def test_folder_named_like_a_temporary_file_still_covers_unknown_descendants(self):
        a = copy.deepcopy(self.complete)
        a["files"] = [
            {
                "path": str(self.root / "scratch.tmp" / "doc.pdf"),
                "parts": ["scratch.tmp", "doc.pdf"],
                "size": 1,
                "sha256": "a" * 64,
            }
        ]
        b = {
            "root": str(self.root),
            "files": [],
            "read_errors": [
                app.read_issue(
                    str(self.root / "scratch.tmp"),
                    ["scratch.tmp"],
                    "directory",
                    PermissionError(13, "Denied"),
                )
            ],
        }
        self.assertEqual(app.compare_snapshots(a, b)[0]["Status"], "Unknown")
