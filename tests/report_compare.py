"""Whole-report comparison; normalize access and generated timestamps only."""

import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def normalized_workbook(path):
    """Compare every ZIP member, cell type, style, formula and layout.

    The only omissions are access-time values, scan start/end timestamps,
    and workbook creation/modification timestamps. Independent scans may
    legitimately produce different values for these fields.
    """
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        headers = list(
            next(
                workbook["folder_review"].iter_rows(
                    min_row=6, max_row=6, values_only=True
                )
            )
        )
        access_column = get_column_letter(headers.index("Accessed at") + 1)
        status = list(workbook["scan_status"].iter_rows(values_only=True))
        volatile_status_rows = {
            index
            for index, row in enumerate(status, 1)
            if row and row[0] in ("Started at", "Completed at")
        }
        status_values = {}
        for row in status[1:]:
            if not row or row[0] is None:
                break
            status_values[row[0]] = row[1]
        rows = sum(1 for _ in workbook["folder_review"].iter_rows(min_row=7))
    finally:
        workbook.close()
    result = {}
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            data = archive.read(name)
            if name in (
                "xl/worksheets/sheet1.xml",
                "xl/worksheets/sheet2.xml",
                "docProps/core.xml",
            ):
                tree = ET.fromstring(data)
                if name == "docProps/core.xml":
                    for element in tree:
                        if element.tag.rsplit("}", 1)[-1] in ("created", "modified"):
                            element.text = "VOLATILE"
                else:
                    for row in tree.findall("s:sheetData/s:row", NS):
                        number = int(row.attrib["r"])
                        for cell in row.findall("s:c", NS):
                            if (
                                name.endswith("sheet1.xml")
                                and number >= 7
                                and cell.attrib["r"] == f"{access_column}{number}"
                            ) or (
                                name.endswith("sheet2.xml")
                                and number in volatile_status_rows
                                and cell.attrib["r"] == f"B{number}"
                            ):
                                for value in cell:
                                    value.text = "VOLATILE"
                data = ET.tostring(tree)
            result[name] = data
    return result, status_values, rows
