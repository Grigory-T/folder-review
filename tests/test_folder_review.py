from __future__ import annotations

import contextlib
import io
import os
import re
import struct
import subprocess
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

import py7zr
from blake3 import blake3
from openpyxl import load_workbook

import folder_review


class FolderReviewTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory(prefix="folder-review-test-")
        self.root = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def make_zip(self, path: Path, members: dict[str, bytes | Path]) -> Path:
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for archive_name, source in members.items():
                if isinstance(source, Path):
                    archive.write(source, archive_name)
                else:
                    archive.writestr(archive_name, source)
        return path

    def make_7z(self, path: Path, members: dict[str, bytes | Path]) -> Path:
        source_dir = self.root / f"{path.stem}-sources"
        source_dir.mkdir(exist_ok=True)
        with py7zr.SevenZipFile(path, "w") as archive:
            for index, (archive_name, source) in enumerate(members.items()):
                if isinstance(source, bytes):
                    data = source
                    source = source_dir / f"member-{index}.bin"
                    source.write_bytes(data)
                archive.write(source, archive_name)
        return path

    def inspect(self, path: Path) -> list[folder_review.ResultRow]:
        stat_result = path.stat()
        item = folder_review.PhysicalFile(
            path=str(path),
            levels=(path.name,),
            size=stat_result.st_size,
            properties=folder_review.file_properties(path.name, stat_result),
        )
        return folder_review.process_file(item)

    def test_virtual_reason_markers_are_filterable_and_last_is_always_one(self) -> None:
        rows = [
            folder_review.folder_placeholder(
                "data/empty", ("empty",), in_archive=False
            ),
            folder_review.folder_placeholder(
                "data/denied",
                ("denied",),
                in_archive=False,
                unreadable=True,
                folder_error="PermissionError: Access is denied",
            ),
        ]
        for name, error in (
            ("encrypted.7z", "PasswordRequired: password is required"),
            ("broken.txt", "CrcError: CRC mismatch"),
            ("large.zip", "BudgetExceeded: row budget exceeded"),
            ("gone.txt", "FileNotFoundError: unavailable"),
        ):
            rows.append(
                folder_review.make_row(
                    "data/" + name, (name,), 123, "Not inspected", False, error
                )
            )
        output = self.root / "markers.xlsx"
        folder_review.write_workbook(rows, output, "data")
        book = load_workbook(output, read_only=True)
        try:
            sheet = book["folder_review"]
            headings = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
            records = [
                dict(zip(headings, row))
                for row in sheet.iter_rows(min_row=7, values_only=True)
            ]
            self.assertEqual(
                [row["File name"] for row in records],
                [
                    "dummy.dummy",
                    "deny.deny",
                    "password.password",
                    "corrupt.corrupt",
                    "limit.limit",
                    "error.error",
                ],
            )
            self.assertTrue(
                all(row["last"] == 1 and row["Virtual row"] == "Yes" for row in records)
            )
            self.assertTrue(all(row["Duplicate group"] is None for row in records))
            self.assertTrue(
                all(row["Type group"] == "Virtual rows" for row in records)
            )
            self.assertEqual(
                [row["Relative path"] for row in records],
                [folder_review.report_relative_path(row.full_path, "data") for row in rows],
            )
            self.assertEqual(records[2]["Level 1"], "encrypted.7z")
            self.assertEqual(records[2]["Level 2"], "password.password")
            self.assertNotIn("Content hash (BLAKE3)", headings)
        finally:
            book.close()
        self.assertEqual(
            rows[2].leaf_name, "encrypted.7z", "Rendering must not mutate scanner rows"
        )

    def test_content_group_ids_are_consecutive_and_shared_only_by_equal_content(
        self,
    ) -> None:
        payloads = [b"alpha", b"beta", b"alpha", b"", None, b"beta", b"unique"]
        rows = [
            replace(
                folder_review.make_row(
                    "data/" + name, (name,), len(data or b""), "Text", False
                ),
                content_hash=blake3(data).digest() if data is not None else None,
            )
            for name, data in zip(
                (
                    "one.txt",
                    "two.txt",
                    "deny.deny",
                    "empty",
                    "unhashed",
                    "three.txt",
                    "unique.txt",
                ),
                payloads,
            )
        ]
        output = self.root / "classes.xlsx"
        folder_review.write_workbook(rows, output, "data", full_metadata=True)
        book = load_workbook(output, read_only=True)
        try:
            sheet = book["folder_review"]
            headings = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
            records = [
                dict(zip(headings, row))
                for row in sheet.iter_rows(min_row=7, values_only=True)
            ]
            self.assertEqual(
                [row["Duplicate group"] for row in records],
                [1, 2, 1, None, None, 2, None],
            )
            self.assertTrue(
                all(row["last"] == 1 and row["Virtual row"] == "No" for row in records)
            )
            self.assertEqual(
                records[2]["File name"],
                "deny.deny",
                "A real marker-like name must remain an ordinary file",
            )
            self.assertEqual(
                records[3]["Content hash (BLAKE3)"],
                "af1349b9f5f9a1a6a0404dea36dcc9499bcb25c9adc112b7cc9a93cae41f3262",
            )
            status = {
                row[0]: row[1]
                for row in book["scan_status"].iter_rows(values_only=True)
                if row[0] is not None
            }
            self.assertEqual(status["Duplicate groups"], 2)
            self.assertEqual(status["Duplicate file rows"], 4)
        finally:
            book.close()

    def test_english_schema_order_units_and_status(self) -> None:
        source = self.root / "пример.txt"
        source.write_bytes(b"x" * 1048576)
        rows = folder_review.inspect_files(
            folder_review.scan_tree(str(self.root), set(), 1)[0], 1
        )
        for full in (False, True):
            output = self.root / ("full.xlsx" if full else "compact.xlsx")
            folder_review.write_workbook(rows, output, "пример", full_metadata=full)
            book = load_workbook(output, read_only=True)
            try:
                sheet = book["folder_review"]
                headers = list(
                    next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
                )[1:]
                self.assertTrue(all(header.isascii() for header in headers))
                self.assertEqual(len(headers), len(set(headers)))
                self.assertEqual(
                    headers.index("File name"),
                    headers.index("Name without extension") + 1,
                )
                self.assertEqual(
                    headers.index("File size (MiB)"),
                    headers.index("File size (bytes)") + 1,
                )
                record = dict(
                    zip(
                        headers,
                        next(sheet.iter_rows(min_row=7, max_row=7, values_only=True))[
                            1:
                        ],
                    )
                )
                self.assertEqual(record["File name"], "пример.txt")
                self.assertEqual(record["Name without extension"], "пример")
                self.assertEqual(record["File size (bytes)"], 1048576)
                self.assertEqual(record["File size (MiB)"], 1)
                self.assertIsNone(record["Duplicate group"])
                self.assertEqual(record["last"], 1)
                status = list(book["scan_status"].iter_rows(values_only=True))
                self.assertEqual(status[0][:2], ("Parameter", "Value"))
                for row in status:
                    self.assertTrue(row[0] is None or row[0].isascii())
                    if row[0] != "Root folder" and isinstance(row[1], str):
                        self.assertTrue(row[1].isascii())
                self.assertIn(("Stage", "Relative path", "Error"), status)
            finally:
                book.close()

    def test_marker_reason_does_not_confuse_names_or_network_errors(self) -> None:
        for error, expected in (
            ("OSError: [WinError 53] Network path not found", "error"),
            ("PermissionError: [WinError 5] Access is denied", "denied"),
            ("FileNotFoundError: missing 'password required.txt'", "error"),
            ("FileNotFoundError: missing 'archive nesting limit.txt'", "error"),
            ("BadZipFile: Bad CRC-32 for file 'password.txt'", "corrupt"),
            ("PasswordRequired: password is required", "password"),
            (
                "RuntimeError: file 'crc.txt' is encrypted, password required for extraction",
                "password",
            ),
            ("BudgetExceeded: expanded-data budget exceeded", "limit"),
        ):
            self.assertEqual(folder_review.problem_kind(error), expected, error)

    def test_file_type_groups_have_two_logical_levels_and_unknown_fallback(
        self,
    ) -> None:
        cases = {
            "book.xlsx": ("Documents", "Spreadsheets"),
            "drawing.vsdx": ("Design and engineering", "Diagrams"),
            "photo.png": ("Images", "Raster images"),
            "film.mp4": ("Video", "Video recordings"),
            "archive.tar.gz": ("Archives and disk images", "Compressed files"),
            "script.py": ("Software and code", "Source code and scripts"),
            "package.whl": ("Software and code", "Packages and installers"),
            "table.csv": ("Data and settings", "Tabular data"),
            "data.json": ("Data and settings", "Structured data"),
            "backup.bak": ("System and support files", "Backups"),
            "plain": ("Other", "No extension"),
            "data.unusual-suffix": ("Other", "Unknown extension"),
            ".env": ("Data and settings", "Settings"),
            "Makefile": ("Software and code", "Build and package metadata"),
        }
        for name, expected in cases.items():
            row = folder_review.make_row(
                "data/" + name, (name,), 0, "Unknown binary", False
            )
            self.assertEqual(folder_review.file_type_groups(row), expected, name)
        extensions = [
            extension
            for _, _, items in folder_review._EXTENSION_GROUPS
            for extension in items.split()
        ]
        self.assertEqual(
            len(extensions), len(set(extensions)), "Each suffix must have one grouping"
        )

    def test_expanded_grouping_and_exact_filename_precedence(self) -> None:
        cases = {
            "CMakeLists.txt": ("Software and code", "Build and package metadata"),
            "meson.build": ("Software and code", "Build and package metadata"),
            "MANIFEST.in": ("Software and code", "Build and package metadata"),
            "pyproject.toml": ("Software and code", "Build and package metadata"),
            "settings.toml": ("Data and settings", "Settings"),
            "py.typed": ("Software and code", "Type declarations"),
            "index.d.mts": ("Software and code", "Type declarations"),
            "index.d.cts": ("Software and code", "Type declarations"),
            "index.d.ts": ("Software and code", "Type declarations"),
            "file.pyi": ("Software and code", "Type declarations"),
            "app.js.map": ("Software and code", "Debug information and source maps"),
            "style.css.map": ("Software and code", "Debug information and source maps"),
            "readme.js.map": ("Software and code", "Debug information and source maps"),
            "license.css.map": ("Software and code", "Debug information and source maps"),
            "module.cts": ("Software and code", "Source code and scripts"),
            "module.tcl": ("Software and code", "Source code and scripts"),
            "module.pyx": ("Software and code", "Source code and scripts"),
            "module.pxd": ("Software and code", "Source code and scripts"),
            "code.class": ("Software and code", "Programs and libraries"),
            "code.pyc": ("Software and code", "Programs and libraries"),
            "code.lib": ("Software and code", "Programs and libraries"),
            "addon.node": ("Software and code", "Programs and libraries"),
            "app.jar": ("Software and code", "Packages and installers"),
            "app.war": ("Software and code", "Packages and installers"),
            "index.j2": ("Software and code", "Templates"),
            "locale.mo": ("Software and code", "Localization and translations"),
            "style.mplstyle": ("Data and settings", "Settings"),
            "data.msgpack": ("Data and settings", "Structured data"),
            "data.npz": ("Data and settings", "Scientific arrays and datasets"),
            "sky.fits": ("Data and settings", "Scientific arrays and datasets"),
            "network.pcapng": ("Data and settings", "Network captures"),
            "weights.tflite": ("Data and settings", "Machine learning models"),
            "assets.pak": ("Data and settings", "Application resources"),
            "data.dat": ("Data and settings", "Unspecified data"),
            "secret.enc": ("Data and settings", "Unspecified data"),
            "LICENSE.python": ("Documents", "Licenses and notices"),
            "README.md": ("Documents", "Text and markup"),
            "server.log": ("System and support files", "Logs"),
            "SHA256SUMS.sha256": (
                "System and support files",
                "Checksums and signatures",
            ),
            ".lock": ("System and support files", "Lock and state files"),
        }
        for name, expected in cases.items():
            with self.subTest(name=name):
                row = folder_review.make_row(name, (name,), 0, "Unknown binary", False)
                self.assertEqual(folder_review.file_type_groups(row), expected)
        for extension in ("obj", "key", "pot", "mts", "map", "pdb"):
            row = folder_review.make_row(
                "file." + extension, ("file." + extension,), 0, "Unknown binary", False
            )
            self.assertEqual(
                folder_review.file_type_groups(row), ("Other", "Ambiguous format")
            )

    def grouping_payloads(self):
        obj = struct.pack("<HHIIIHH", 0x8664, 1, 0, 0, 0, 0, 0) + bytes(40)
        ts = bytes([0x47]) + bytes(187)
        return {
            "test.key": (
                b"-----BEGIN PRIVATE KEY-----\nnot-a-real-key\n-----END PRIVATE KEY-----\n",
                ("System and support files", "Certificates and keys"),
            ),
            "messages.pot": (
                b'# translation template\nmsgid ""\nmsgstr ""\n',
                ("Software and code", "Localization and translations"),
            ),
            "code.obj": (obj, ("Software and code", "Programs and libraries")),
            "mesh.obj": (
                b"v 0 0 0\nv 1 1 1\nf 1 2 3\n",
                ("Design and engineering", "3D models"),
            ),
            "index.mts": (
                b"export const value = 1;\n",
                ("Software and code", "Source code and scripts"),
            ),
            "film.mts": (ts * 3, ("Video", "Video recordings")),
            "film.ts": (ts * 3, ("Video", "Video recordings")),
            "debug.pdb": (
                b"Microsoft C/C++ MSF 7.00\r\n\x1aDS\x00\x00\x00" + bytes(64),
                ("Software and code", "Debug information and source maps"),
            ),
            "native.bare": (
                b"\x7fELF" + bytes(64),
                ("Software and code", "Programs and libraries"),
            ),
            "native.unknown": (
                b"\xcf\xfa\xed\xfe" + bytes(64),
                ("Software and code", "Programs and libraries"),
            ),
            "CMakeLists.txt": (
                b"project(example)\n",
                ("Software and code", "Build and package metadata"),
            ),
            "app.js.map": (
                b'{"version":3,"sources":[]}',
                ("Software and code", "Debug information and source maps"),
            ),
            "license.js.map": (
                b'{"version":3,"sources":[]}',
                ("Software and code", "Debug information and source maps"),
            ),
        }

    def test_grouping_content_evidence_matches_disk_zip_and_7z(self) -> None:
        data = self.grouping_payloads()
        inputs = self.root / "group-input"
        inputs.mkdir()
        for name, (payload, _) in data.items():
            (inputs / name).write_bytes(payload)
        zip_path = self.make_zip(
            self.root / "groups.zip",
            {name: payload for name, (payload, _) in data.items()},
        )
        seven = self.make_7z(
            self.root / "groups.7z",
            {name: payload for name, (payload, _) in data.items()},
        )
        disk = folder_review.inspect_files(
            folder_review.scan_tree(str(inputs), set(), 2)[0], 2
        )
        for rows in (disk, self.inspect(zip_path), self.inspect(seven)):
            self.assertEqual(len(rows), len(data))
            for row in rows:
                self.assertIsNone(row.error)
                self.assertEqual(
                    folder_review.file_type_groups(row),
                    data[row.leaf_name][1],
                    row.leaf_name,
                )
                self.assertEqual(
                    row.content_hash, blake3(data[row.leaf_name][0]).digest()
                )

    def test_grouping_hints_preserve_archive_routing_and_reject_short_headers(
        self,
    ) -> None:
        for signature, name in (
            (b"PK\x03\x04", "ZIP container"),
            (b"Rar!\x1a\x07\x00", "RAR archive"),
            (b"7z\xbc\xaf\x27\x1c", "7z archive"),
        ):
            self.assertEqual(
                folder_review.refine_grouping_content("data.zip", signature, name), name
            )
        for name, data in (
            ("bad.obj", b"\x64\x86"),
            ("bad.pdb", b"Microsoft C/C++"),
            ("bad.mts", b"G"),
            ("bad.pot", b"msgid only"),
            ("bad.key", b"private key"),
        ):
            content = folder_review.member_content_type(data, name)
            row = folder_review.make_row(name, (name,), len(data), content, False)
            expected = (
                ("Software and code", "Source code and scripts")
                if name.endswith(".mts")
                else ("Other", "Ambiguous format")
            )
            self.assertEqual(folder_review.file_type_groups(row), expected)
        row = folder_review.make_row(
            "slides.pot", ("slides.pot",), 0, "PowerPoint presentation (PPT)", False
        )
        self.assertEqual(
            folder_review.file_type_groups(row), ("Documents", "Presentations")
        )
        row = folder_review.make_row(
            "slides.key", ("slides.key",), 0, "ZIP container (leaf file)", False
        )
        self.assertEqual(
            folder_review.file_type_groups(row), ("Documents", "Presentations")
        )

    def test_fingerprints_match_physical_zip_7z_and_nested_content(self) -> None:
        payload = b"\xef\xbb\xbfduplicate payload" + bytes(range(256)) * 5000
        expected = blake3(payload).digest()
        first = self.root / "first.bin"
        first.write_bytes(payload)
        second = self.root / "different-name.txt"
        second.write_bytes(payload)
        zipped = self.make_zip(
            self.root / "data.zip", {"payload.bin": payload, "empty": b""}
        )
        seven = self.make_7z(
            self.root / "data.7z", {"payload.bin": payload, "empty": b""}
        )
        nested = self.make_zip(self.root / "outer.zip", {"inner.7z": seven})
        for path in (first, second, zipped, seven, nested):
            with self.subTest(path=path.name):
                rows = self.inspect(path)
                self.assertTrue(all(row.error is None for row in rows), rows)
                for row in rows:
                    self.assertEqual(
                        row.content_hash,
                        blake3(b"").digest() if row.leaf_name == "empty" else expected,
                    )
        self.assertEqual(
            self.inspect(first)[0].content_hash, self.inspect(second)[0].content_hash
        )

    def test_incomplete_or_changed_content_has_no_hash(self) -> None:
        source = self.root / "data.txt"
        source.write_bytes(b"data")
        item = folder_review.PhysicalFile(str(source), (source.name,), 5)
        row = folder_review.process_file(item)[0]
        self.assertIsNone(row.content_hash)
        self.assertIn("changed during inspection", row.error)
        with folder_review.archive_temp_directory() as directory:
            row = folder_review.inspect_member_stream(
                io.BytesIO(b"partial"),
                size=100,
                logical_path="data.zip::member.txt",
                levels=("data.zip", "member.txt"),
                parent_depth=0,
                temp_dir=directory,
                budget=folder_review.WorkBudget(),
            )[0]
        self.assertIsNone(row.content_hash)
        self.assertIsNotNone(row.error)
        self.assertEqual(row.integrity_status, folder_review.INTEGRITY_ERROR)
        before = source.stat()
        after = SimpleNamespace(
            **{
                key: getattr(before, key)
                for key in (
                    "st_size",
                    "st_mtime_ns",
                    "st_ctime_ns",
                    "st_ino",
                    "st_dev",
                )
            }
        )
        after.st_mtime_ns += 1
        with (
            mock.patch.object(folder_review.os, "fstat", side_effect=[before, after]),
            self.assertRaisesRegex(OSError, "changed during inspection"),
        ):
            folder_review.file_digest(str(source))

    def test_empty_folders_have_minimal_virtual_markers_and_zero_file_count(
        self,
    ) -> None:
        root = self.root / "input"
        (root / "chain" / "empty").mkdir(parents=True)
        (root / "sibling").mkdir()
        (root / "filled").mkdir()
        (root / "filled" / "dummy.dummy").write_bytes(b"real file")
        items, _stats = folder_review.scan_tree(str(root), set(), 2)
        rows = folder_review.inspect_files(items, 2)
        virtual = [row for row in rows if row.virtual_folder]
        self.assertEqual(len(virtual), 2)
        self.assertEqual(
            {row.levels for row in virtual},
            {("chain", "empty", "dummy.dummy"), ("sibling", "dummy.dummy")},
        )
        self.assertTrue(
            all(row.content_hash is None and row.extension == "" for row in virtual)
        )
        self.assertEqual(sum(not row.virtual_folder for row in rows), 1)
        self.assertFalse((root / "sibling" / "dummy.dummy").exists())
        self.assertFalse((root / "chain" / "empty" / "dummy.dummy").exists())
        output = self.root / "folders.xlsx"
        folder_review.write_workbook(rows, output, str(root))
        book = load_workbook(output, read_only=True)
        try:
            sheet = book["folder_review"]
            headers = list(
                next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
            )
            values = list(sheet.iter_rows(min_row=7, values_only=True))
            self.assertEqual(sum(row[headers.index("last")] for row in values), 3)
            self.assertEqual(
                sum(
                    row[headers.index("Duplicate group")] is not None for row in values
                ),
                0,
            )
        finally:
            book.close()
        empty = self.root / "empty-root"
        empty.mkdir()
        rows = folder_review.inspect_files(
            folder_review.scan_tree(str(empty), set(), 1)[0], 1
        )
        self.assertEqual(len(rows), 1)
        self.assertTrue(rows[0].virtual_folder)

    def test_archive_empty_directories_are_retained_without_staging(self) -> None:
        zipped = self.make_zip(
            self.root / "folders.zip",
            {
                "parent/": b"",
                "parent/empty/": b"",
                "other/": b"",
                "filled/": b"",
                "filled/file.txt": b"content",
            },
        )
        source = self.root / "directories"
        (source / "parent" / "empty").mkdir(parents=True)
        (source / "other").mkdir()
        (source / "filled").mkdir()
        (source / "filled" / "file.txt").write_bytes(b"content")
        seven = self.root / "folders.7z"
        with py7zr.SevenZipFile(seven, "w") as archive:
            archive.writeall(source, arcname="collection")
        for path in (zipped, seven):
            rows = self.inspect(path)
            markers = [row for row in rows if row.virtual_folder]
            self.assertEqual(len(markers), 2, rows)
            self.assertEqual({row.levels[-2] for row in markers}, {"empty", "other"})
            self.assertTrue(
                all(
                    row.content_hash is None and row.in_archive == "Yes"
                    for row in markers
                )
            )
            self.assertTrue(all(row.error is None for row in rows), rows)

    def test_rar_folder_markers_and_stream_hash(self) -> None:
        source = self.root / "mock.rar"
        source.write_bytes(b"mock archive")
        node = folder_review.ArchiveNode(
            str(source), str(source), (source.name,), source.stat().st_size, 0
        )
        payload = b"RAR member content"
        archive = mock.MagicMock()
        archive.__enter__.return_value = archive
        archive.infolist.return_value = [
            SimpleNamespace(filename="empty/deep/", is_dir=lambda: True),
            SimpleNamespace(
                filename="file.txt", file_size=len(payload), is_dir=lambda: False
            ),
        ]
        archive.open.side_effect = lambda info: io.BytesIO(payload)
        with (
            mock.patch.object(folder_review.rarfile, "RarFile", return_value=archive),
            folder_review.archive_temp_directory() as directory,
        ):
            rows = folder_review.expand_rar(
                node, directory, folder_review.WorkBudget(rows_claimed=1)
            )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].content_hash, blake3(payload).digest())
        self.assertTrue(rows[1].virtual_folder)
        self.assertEqual(rows[1].levels, ("mock.rar", "empty", "deep", "dummy.dummy"))

    def test_folder_markers_obey_logical_row_limit(self) -> None:
        root = self.root / "input"
        for name in ("a", "b", "c"):
            (root / name).mkdir(parents=True)
        items, stats = folder_review.scan_tree(str(root), set(), 1)
        selected, physical_count = folder_review.apply_physical_row_limit(
            items, 2, stats, str(root)
        )
        self.assertEqual(physical_count, 0)
        self.assertEqual(len(selected), 2)
        self.assertEqual(stats.skipped_entries, 1)
        archive = self.make_zip(self.root / "only-folders.zip", {"empty/": b""})
        row = folder_review.process_file(
            folder_review.PhysicalFile(
                str(archive), (archive.name,), archive.stat().st_size
            ),
            folder_review.WorkBudget(max_rows=1, rows_claimed=1),
        )[0]
        self.assertIn("Folder placeholders omitted", row.error)

    def test_extension_heuristics_preserve_real_suffixes_and_raw_metadata(self) -> None:
        expected = {
            "1.3.3-cp314-cp314-win_amd64": "",
            "3.1.5-py2.py3-none-any": "",
            "cpython-3.14-windows-x86_64-none": "",
            "1.2.3": "",
            "cache.32_13426590566142933": "",
            "cache.4_13429547779578144": "",
            "ttx.1": "1",
            "openssl.3": "3",
            "archive.tar.gz": "gz",
            "name.jupyterlab-workspace": "jupyterlab-workspace",
            "template.nuspectemplate": "nuspectemplate",
            "data.db-journal": "db-journal",
            "package-1.3.3-cp314-cp314-win_amd64.whl": "whl",
            "report.vsdx": "vsdx",
            "1.2.3-release.zip": "zip",
            "1.2.3-py3-none-any.whl": "whl",
        }
        for name, suffix in expected.items():
            self.assertEqual(folder_review.extension_of(name), suffix, name)
        row = folder_review.make_row("data/1.2.3", ("1.2.3",), 0, "Text", False)
        output = self.root / "suffix.xlsx"
        folder_review.write_workbook([row], output, "data", full_metadata=True)
        book = load_workbook(output, read_only=True)
        try:
            sheet = book["folder_review"]
            headers = list(
                next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
            )
            values = next(sheet.iter_rows(min_row=7, values_only=True))
            self.assertEqual(values[headers.index("Name without extension")], "1.2.3")
            self.assertIn(values[headers.index("Extension")], (None, ""))
            self.assertEqual(values[headers.index("Raw suffix")], "3")
        finally:
            book.close()

    def test_writer_text_types_unicode_and_dates(self) -> None:
        source = self.root / "metadata.txt"
        source.write_text("sample", encoding="utf-8")
        properties = replace(
            folder_review.file_properties(source.name, source.stat()),
            inode="18446744073709551615",
            modified_at=datetime(2025, 1, 2, 3, 4, 5),
        )
        names = [
            "=1+1.txt",
            "+formula.txt",
            "@name.txt",
            "https://example.invalid",
            "00123.txt",
            "пример-😀.txt",
            "literal_x000D_.txt",
            "line\rbreak.txt",
            "bad\x01control.txt",
            "long-" + "x" * 32770,
        ]
        rows = [
            folder_review.make_row(
                "\\\\server\\share\\" + name,
                (name,),
                1234,
                "Text",
                False,
                properties=properties,
            )
            for name in names
        ]
        output = self.root / "text.xlsx"
        folder_review.write_workbook(rows, output, "sample", full_metadata=True)
        book = load_workbook(output, read_only=True, data_only=False)
        try:
            ws = book["folder_review"]
            headers = list(next(ws.iter_rows(min_row=6, max_row=6, values_only=True)))
            name_column = headers.index("File name") + 1
            for row in ws.iter_rows(min_row=7):
                self.assertEqual(row[name_column - 1].data_type, "s")
                self.assertEqual(
                    row[headers.index("File ID (st_ino)")].value, properties.inode
                )
                self.assertEqual(
                    row[headers.index("Modified at")].value,
                    properties.modified_at,
                )
                self.assertEqual(row[headers.index("File size (bytes)")].data_type, "n")
        finally:
            book.close()
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with zipfile.ZipFile(output) as archive:
            self.assertIsNone(archive.testzip())
            tree = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        self.assertEqual(tree.findall(".//s:f", ns), [])
        self.assertEqual(tree.findall(".//s:hyperlink", ns), [])
        for number, name in enumerate(names, 7):
            coordinate = f"{folder_review.get_column_letter(name_column)}{number}"
            cell = tree.find(f".//s:c[@r='{coordinate}']/s:is", ns)
            encoded = "".join(cell.itertext())
            # Decode Excel's escape convention once; escaped literal sequences
            # remain literal. openpyxl 3.1.5 does not decode inline strings.
            decoded = re.sub(
                r"_x([0-9a-fA-F]{4})_", lambda match: chr(int(match[1], 16)), encoded
            )
            self.assertEqual(decoded, folder_review.excel_text(name))

    def test_writer_empty_report_and_streaming_row_heights(self) -> None:
        output = self.root / "empty.xlsx"
        folder_review.write_workbook([], output, "sample")
        book = load_workbook(output, read_only=False)
        try:
            self.assertEqual(book["folder_review"].max_row, 6)
            self.assertEqual(book["folder_review"].row_dimensions[6].height, 72)
            self.assertEqual(book["scan_status"].row_dimensions[17].height, 32)
            self.assertEqual(book["scan_status"].freeze_panes, "A2")
            self.assertEqual(book["scan_status"]["B2"].value, "COMPLETE")
        finally:
            book.close()

    def test_writer_failure_preserves_report_and_cleans_temporary_files(self) -> None:
        output = self.root / "existing.xlsx"
        output.write_bytes(b"previous report")
        with tempfile.TemporaryDirectory(dir=self.root) as work:
            with (
                mock.patch.object(tempfile, "tempdir", work),
                mock.patch.object(
                    folder_review._ReportSheet,
                    "append",
                    side_effect=OSError("writer failed"),
                ),
                self.assertRaisesRegex(OSError, "writer failed"),
            ):
                folder_review.write_workbook([], output, "sample")
            self.assertEqual(list(Path(work).iterdir()), [])
        self.assertEqual(output.read_bytes(), b"previous report")
        self.assertEqual(list(self.root.glob(".folder-review-*.tmp.xlsx")), [])

    def test_zip_7z_zip_and_office_package(self) -> None:
        office = self.make_zip(
            self.root / "book.xlsx",
            {
                "[Content_Types].xml": b"<Types />",
                "_rels/.rels": b"<Relationships />",
                "xl/workbook.xml": b"<workbook />",
            },
        )
        inner = self.make_zip(
            self.root / "inner.zip",
            {
                "Deep/final.pdf": b"%PDF-1.4\n%%EOF\n",
                "Deep/book.xlsx": office,
                "Deep/bom.txt": b"\xff\xfeh\x00i\x00",
            },
        )
        middle = self.make_7z(
            self.root / "middle.7z",
            {
                "Layer2/inner.zip": inner,
                "Layer2/note.txt": b"terminal text",
            },
        )
        outer = self.make_zip(
            self.root / "outer.zip",
            {
                "Layer1/middle.7z": middle,
                "Direct/picture.bmp": b"BM" + bytes(32),
            },
        )

        rows = self.inspect(outer)
        paths = {row.full_path for row in rows}

        self.assertTrue(
            any(
                path.endswith(
                    r"outer.zip::Layer1\middle.7z::Layer2\inner.zip::Deep\final.pdf"
                )
                for path in paths
            )
        )
        workbook_row = next(row for row in rows if row.leaf_name == "book.xlsx")
        bom_row = next(row for row in rows if row.leaf_name == "bom.txt")
        self.assertEqual(workbook_row.content_type, "Excel workbook (XLSX)")
        self.assertEqual(bom_row.bom, "UTF-16 LE")
        self.assertEqual(bom_row.integrity_status, folder_review.INTEGRITY_VERIFIED)
        self.assertEqual(bom_row.item_type, "Archive member")
        self.assertIsNone(bom_row.properties)
        self.assertFalse(
            any(row.leaf_name in {"middle.7z", "inner.zip"} for row in rows)
        )
        self.assertFalse(any(row.error for row in rows))

    def test_7z_zip_7z_reverse_order(self) -> None:
        inner = self.make_7z(
            self.root / "inner.7z",
            {"Bottom/leaf.txt": b"leaf"},
        )
        middle = self.make_zip(
            self.root / "middle.zip",
            {"Middle/inner.7z": inner},
        )
        outer = self.make_7z(
            self.root / "outer.7z",
            {"Top/middle.zip": middle},
        )

        rows = self.inspect(outer)

        self.assertEqual(len(rows), 1)
        self.assertTrue(
            rows[0].full_path.endswith(
                r"outer.7z::Top\middle.zip::Middle\inner.7z::Bottom\leaf.txt"
            )
        )
        self.assertEqual(rows[0].content_type, "Text (UTF-8 / ASCII)")
        self.assertIsNone(rows[0].error)

    def test_opaque_suffix_does_not_expand_archive_content(self) -> None:
        archive = self.make_zip(
            self.root / "archive.bin",
            {"Folder/leaf.txt": b"leaf"},
        )

        rows = self.inspect(archive)

        self.assertEqual([row.leaf_name for row in rows], ["archive.bin"])
        self.assertEqual(rows[0].content_type, "ZIP container (leaf file)")
        self.assertFalse(rows[0].in_archive == "Yes")
        self.assertIsNone(rows[0].error)

    def test_ordinary_zip_with_office_like_folder_is_expanded(self) -> None:
        archive = self.make_zip(
            self.root / "ordinary.ZIP",
            {
                "word/notes.txt": b"not an Office package",
                "visio/document.xml": b"<not-a-package/>",
                "other.txt": b"must also be inventoried",
            },
        )

        rows = self.inspect(archive)

        self.assertEqual(
            {row.leaf_name for row in rows},
            {"notes.txt", "document.xml", "other.txt"},
        )
        self.assertTrue(all(row.error is None for row in rows))
        self.assertTrue(
            all(
                row.integrity_status == folder_review.INTEGRITY_VERIFIED for row in rows
            )
        )

    def test_zip_based_leaf_formats_remain_one_row_everywhere(self) -> None:
        package = self.make_zip(
            self.root / "package.zip",
            {
                "[Content_Types].xml": b"<Types/>",
                "_rels/.rels": b"<Relationships/>",
                "visio/document.xml": b"<VisioDocument/>",
                "visio/masters/_rels/master43.xml.rels": b"<Relationships/>",
                "visio/media/image2.emf": bytes(100),
                "embedded/hidden.zip": b"PK\x03\x04bad",
            },
        )
        formats = {*folder_review.ZIP_LEAF_FORMATS, "unknown-format"}
        members = {}
        expected_types = {}
        for extension in sorted(formats):
            name = f"document.{extension.upper()}"
            path = self.root / name
            path.write_bytes(package.read_bytes())
            rows = self.inspect(path)
            self.assertEqual(len(rows), 1, name)
            self.assertEqual(rows[0].leaf_name, name)
            self.assertIsNone(rows[0].error, name)
            expected_types[name] = rows[0].content_type
            members[name] = path
        for suffix, create in (("zip", self.make_zip), ("7z", self.make_7z)):
            outer = create(self.root / f"outer.{suffix.upper()}", members)
            with mock.patch.object(
                folder_review, "make_temp_path", wraps=folder_review.make_temp_path
            ) as stage:
                rows = self.inspect(outer)
            self.assertEqual({row.leaf_name for row in rows}, set(members))
            self.assertEqual(len(rows), len(members))
            for row in rows:
                self.assertIsNone(row.error, row.leaf_name)
                self.assertEqual(row.integrity_status, folder_review.INTEGRITY_VERIFIED)
                if row.extension in {"vsdx", "vsdm", "vstx", "vstm", "vssx", "vssm"}:
                    self.assertEqual(row.content_type, expected_types[row.leaf_name])
                self.assertEqual(row.full_path.count("::"), 1)
            stage.assert_not_called()

    def test_packaged_documents_renamed_zip_are_still_leaves(self) -> None:
        for name, parts in (
            ("visio", {"visio/document.xml": b"<VisioDocument/>"}),
            (
                "open-document",
                {
                    "mimetype": b"application/vnd.oasis.opendocument.text",
                    "META-INF/manifest.xml": b"<manifest/>",
                },
            ),
            ("other-opc", {"FixedDocumentSequence.fdseq": b"<FixedDocumentSequence/>"}),
        ):
            members = {
                "[Content_Types].xml": b"<Types/>",
                "_rels/.rels": b"<Relationships/>",
                **parts,
            }
            path = self.make_zip(self.root / f"{name}.zip", members)
            outer = self.make_7z(self.root / f"{name}-outer.7z", {path.name: path})
            for source in (path, outer):
                rows = self.inspect(source)
                self.assertEqual([row.leaf_name for row in rows], [path.name])
                self.assertIsNone(rows[0].error)

    def test_known_package_content_takes_priority_over_misleading_suffix(self) -> None:
        for suffix, part, expected in (
            ("xlsx", "word/document.xml", "Word document (DOCX)"),
            ("docx", "xl/workbook.xml", "Excel workbook (XLSX)"),
            ("xlsx", "visio/document.xml", "Visio document (ZIP package)"),
        ):
            path = self.make_zip(
                self.root / f"misnamed-{part.split('/')[0]}.{suffix}",
                {
                    "[Content_Types].xml": b"<Types/>",
                    "_rels/.rels": b"<Relationships/>",
                    part: b"<document/>",
                },
            )
            rows = self.inspect(path)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].content_type, expected)
            self.assertIsNone(rows[0].error)

    def test_archive_suffix_does_not_override_file_content(self) -> None:
        for suffix in ("zip", "rar", "7z"):
            path = self.root / f"text.{suffix}"
            path.write_bytes(b"ordinary text")
            rows = self.inspect(path)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].content_type, "Text (UTF-8 / ASCII)")
            self.assertIsNone(rows[0].error)

    def test_rar_leaf_packages_remain_leaves(self) -> None:
        fixture = Path(__file__).parent / "fixtures" / "windows" / "leaf-packages.rar"
        archive = self.root / "leaf-packages.rar"
        archive.write_bytes(fixture.read_bytes())
        expected = {
            f"document.{extension.upper()}"
            for extension in folder_review.ZIP_LEAF_FORMATS
        }
        expected.update({"document.UNKNOWN-FORMAT", "document.zip"})
        for source in (
            archive,
            self.make_zip(self.root / "wrapped.zip", {archive.name: archive}),
            self.make_7z(self.root / "wrapped.7z", {archive.name: archive}),
        ):
            rows = self.inspect(source)
            self.assertEqual({row.leaf_name for row in rows}, expected)
            self.assertEqual(len(rows), len(expected))
            self.assertFalse(
                any(row.error for row in rows), [row.error for row in rows if row.error]
            )
            self.assertTrue(
                all(
                    row.integrity_status == folder_review.INTEGRITY_VERIFIED
                    for row in rows
                )
            )
        comic = self.root / "comic.cbr"
        comic.write_bytes(fixture.read_bytes())
        rows = self.inspect(comic)
        self.assertEqual([row.leaf_name for row in rows], ["comic.cbr"])
        self.assertIsNone(rows[0].error)

    def test_leaf_retains_containing_zip_crc_errors(self) -> None:
        archive = self.make_zip(
            self.root / "bad-parent.zip", {"diagram.vsdx": b"PK\x03\x04" + b"A" * 65536}
        )
        data = bytearray(archive.read_bytes())
        central = data.index(b"PK\x01\x02")
        bad_crc = struct.unpack_from("<I", data, central + 16)[0] ^ 1
        struct.pack_into("<I", data, central + 16, bad_crc)
        struct.pack_into("<I", data, 14, bad_crc)
        archive.write_bytes(data)
        rows = self.inspect(archive)
        self.assertEqual([row.leaf_name for row in rows], ["diagram.vsdx"])
        self.assertIsNotNone(rows[0].error)
        self.assertEqual(rows[0].integrity_status, folder_review.INTEGRITY_ERROR)

    def test_leaf_member_does_not_consume_archive_depth_or_staging_limit(self) -> None:
        import io

        content = b"PK\x03\x04document bytes"
        with (
            folder_review.archive_temp_directory() as directory,
            mock.patch.object(
                folder_review, "MAX_NESTED_ARCHIVE_SIZE", len(content) - 1
            ),
        ):
            rows = folder_review.inspect_member_stream(
                io.BytesIO(content),
                size=len(content),
                logical_path="outer.zip::document.vsdx",
                levels=("outer.zip", "document.vsdx"),
                parent_depth=folder_review.MAX_ARCHIVE_DEPTH - 1,
                temp_dir=directory,
                budget=folder_review.WorkBudget(),
            )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].content_type, "Visio drawing (VSDX)")
        self.assertIsNone(rows[0].error)

    def test_bad_zip_crc_is_reported(self) -> None:
        archive_path = self.root / "bad-crc.zip"
        with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_STORED) as archive:
            archive.writestr("large.txt", b"A" * 65_536)
        data = bytearray(archive_path.read_bytes())
        central = data.index(b"PK\x01\x02")
        original_crc = struct.unpack_from("<I", data, central + 16)[0]
        bad_crc = original_crc ^ 1
        struct.pack_into("<I", data, central + 16, bad_crc)
        struct.pack_into("<I", data, 14, bad_crc)
        archive_path.write_bytes(data)

        rows = self.inspect(archive_path)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].leaf_name, "large.txt")
        self.assertIsNotNone(rows[0].error)
        self.assertIn("BadZipFile", rows[0].error)
        self.assertIsNone(rows[0].content_hash)
        self.assertEqual(rows[0].integrity_status, folder_review.INTEGRITY_ERROR)

    def test_bad_7z_crc_is_reported_after_all_bytes_are_captured(self) -> None:
        archive_path = self.root / "bad-crc.7z"
        with py7zr.SevenZipFile(
            archive_path,
            "w",
            filters=[{"id": py7zr.FILTER_COPY}],
        ) as archive:
            archive.set_encoded_header_mode(False)
            archive.writestr(b"A" * 65_536, "payload.txt")

        data = bytearray(archive_path.read_bytes())
        self.assertEqual(data[32], ord("A"), "Unexpected stored 7z layout")
        data[32] = ord("B")
        archive_path.write_bytes(data)

        rows = self.inspect(archive_path)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].leaf_name, "payload.txt")
        self.assertEqual(rows[0].content_type, "Text (UTF-8 / ASCII)")
        self.assertIsNotNone(rows[0].error)
        self.assertIn("CrcError", rows[0].error)
        self.assertIsNone(rows[0].content_hash)
        self.assertEqual(rows[0].integrity_status, folder_review.INTEGRITY_ERROR)

    def test_encrypted_7z_preserves_known_member_names(self) -> None:
        archive_path = self.root / "encrypted.7z"
        with py7zr.SevenZipFile(
            archive_path,
            "w",
            password="review-only",
            header_encryption=False,
        ) as archive:
            archive.writestr(b"first", "first.txt")
            archive.writestr(b"second", "second.txt")

        rows = self.inspect(archive_path)

        self.assertEqual(
            {row.leaf_name for row in rows},
            {"first.txt", "second.txt"},
        )
        self.assertTrue(all(row.error for row in rows))
        self.assertTrue(all(row.content_hash is None for row in rows))
        self.assertTrue(all("PasswordRequired" in row.error for row in rows))
        self.assertTrue(
            all(
                row.integrity_status == folder_review.INTEGRITY_NOT_VERIFIED
                for row in rows
            )
        )

    def test_encrypted_7z_headers_are_partial(self) -> None:
        archive_path = self.root / "encrypted-headers.7z"
        with py7zr.SevenZipFile(
            archive_path,
            "w",
            password="review-only",
            header_encryption=True,
        ) as archive:
            archive.writestr(b"secret", "secret.txt")

        rows = self.inspect(archive_path)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].leaf_name, "encrypted-headers.7z")
        self.assertEqual(rows[0].content_type, "7z archive")
        self.assertEqual(
            rows[0].integrity_status,
            folder_review.INTEGRITY_NOT_VERIFIED,
        )
        self.assertEqual(
            rows[0].error,
            "Encrypted 7z headers could not be inspected without a password",
        )

    def test_encrypted_rar_headers_produce_partial_workbook(self) -> None:
        input_root = self.root / "input"
        input_root.mkdir()
        fixture = (
            Path(__file__).parent
            / "fixtures"
            / "windows"
            / "rar5-encrypted-headers.rar"
        )
        archive_path = input_root / "encrypted-headers.rar"
        archive_path.write_bytes(fixture.read_bytes())
        output = self.root / "encrypted-headers.xlsx"

        rows = self.inspect(archive_path)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].leaf_name, "encrypted-headers.rar")
        self.assertEqual(rows[0].content_type, "RAR archive")
        self.assertEqual(
            rows[0].integrity_status,
            folder_review.INTEGRITY_NOT_VERIFIED,
        )
        self.assertEqual(
            rows[0].error,
            "Encrypted RAR headers could not be inspected without a password",
        )

        with (
            mock.patch.object(
                folder_review.sys,
                "argv",
                [
                    "folder_review.py",
                    "--root",
                    str(input_root),
                    "--output",
                    str(output),
                ],
            ),
            contextlib.redirect_stdout(io.StringIO()),
            contextlib.redirect_stderr(io.StringIO()),
        ):
            exit_code = folder_review.run()

        self.assertEqual(exit_code, 2)
        workbook = load_workbook(output, read_only=True, data_only=True)
        try:
            status_values = list(workbook["scan_status"].iter_rows(values_only=True))
            self.assertIn(("Status", "PARTIAL"), [row[:2] for row in status_values])
            self.assertIn(("Rows with errors", 1), [row[:2] for row in status_values])
        finally:
            workbook.close()

    def test_genuinely_empty_unencrypted_rar_remains_verified(self) -> None:
        archive_path = self.root / "empty.rar"
        archive_path.write_bytes(b"fixture content is mocked")
        stat_result = archive_path.stat()
        node = folder_review.ArchiveNode(
            str(archive_path),
            str(archive_path),
            (archive_path.name,),
            stat_result.st_size,
            0,
        )
        archive = mock.MagicMock()
        archive.__enter__.return_value = archive
        archive.infolist.return_value = []
        archive.needs_password.return_value = False

        with (
            mock.patch.object(folder_review.rarfile, "RarFile", return_value=archive),
            tempfile.TemporaryDirectory() as temp_dir,
        ):
            rows = folder_review.expand_rar(
                node,
                temp_dir,
                folder_review.WorkBudget(rows_claimed=1),
            )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].content_type, "RAR archive (empty)")
        self.assertEqual(
            rows[0].content_hash, blake3(archive_path.read_bytes()).digest()
        )
        self.assertEqual(rows[0].integrity_status, folder_review.INTEGRITY_VERIFIED)
        self.assertIsNone(rows[0].error)

    def test_unusable_archive_member_path_is_not_silently_dropped(self) -> None:
        archive_path = self.root / "unusable-member-path.zip"
        with zipfile.ZipFile(archive_path, "w") as archive:
            archive.writestr(".", b"payload")

        rows = self.inspect(archive_path)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].content_type, "Not inspected")
        self.assertEqual(rows[0].integrity_status, folder_review.INTEGRITY_ERROR)
        self.assertIn("no usable components", rows[0].error)

    def test_archive_expansion_budget_stops_7z_work(self) -> None:
        archive_path = self.make_7z(
            self.root / "large.7z",
            {"large.txt": b"A" * (2 * 1024 * 1024)},
        )
        stat_result = archive_path.stat()
        item = folder_review.PhysicalFile(
            path=str(archive_path),
            levels=(archive_path.name,),
            size=stat_result.st_size,
            properties=folder_review.file_properties(
                archive_path.name,
                stat_result,
            ),
        )
        budget = folder_review.WorkBudget(
            max_expanded_bytes=1024 * 1024,
            rows_claimed=1,
        )

        rows = folder_review.process_file(item, budget)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].leaf_name, "large.txt")
        self.assertIn("budget", rows[0].error.lower())
        self.assertEqual(rows[0].integrity_status, folder_review.INTEGRITY_NOT_VERIFIED)
        self.assertLessEqual(budget.expanded_bytes, 1024 * 1024)

    def test_corrupt_nested_archive_becomes_error_row(self) -> None:
        outer = self.make_zip(
            self.root / "outer.zip",
            {"Folder/broken.zip": b"PK\x03\x04not-a-valid-zip"},
        )

        rows = self.inspect(outer)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].leaf_name, "broken.zip")
        self.assertEqual(rows[0].content_type, "ZIP container (unreadable)")
        self.assertEqual(rows[0].error, "Invalid or unsupported ZIP container")

    def test_zip_rar_zip_order(self) -> None:
        fixture = Path(__file__).parent / "fixtures" / "rar-with-zip.rar"
        outer = self.make_zip(
            self.root / "outer.zip",
            {"Top/middle.rar": fixture},
        )
        folder_review.configure_rar_backend()

        rows = self.inspect(outer)
        if len(rows) == 1 and rows[0].error == "RarCannotExec":
            self.skipTest("No RAR extraction backend is installed")

        self.assertTrue(
            any(
                row.full_path.endswith(
                    r"outer.zip::Top\middle.rar::Outer\nested.zip::Inner\leaf.txt"
                )
                for row in rows
            )
        )
        self.assertFalse(any(row.error for row in rows))

    def test_workbook_contract(self) -> None:
        source = self.root / "leaf.pdf"
        source.write_bytes(b"%PDF-1.4\n")
        properties = folder_review.file_properties(source.name, source.stat())
        rows = [
            folder_review.make_row(
                str(source),
                ("leaf.pdf",),
                source.stat().st_size,
                "PDF",
                False,
                bom=folder_review.BOM_NOT_PRESENT,
                properties=properties,
            )
        ]
        output = self.root / "folder-review.xlsx"

        folder_review.write_workbook(rows, output, "sample", full_metadata=True)

        workbook = load_workbook(output, read_only=True, data_only=False)
        try:
            sheet = workbook["folder_review"]
            headers = [
                cell.value for cell in next(sheet.iter_rows(min_row=6, max_row=6))
            ]
            values = next(sheet.iter_rows(min_row=7, max_row=7, values_only=True))
            columns = {header: index for index, header in enumerate(headers) if header}
            requested_headers = {
                "Relative path",
                "Name without extension",
                "Extension",
                "File name",
                "Item type",
                "Detected role",
                "Link type",
                "Link target",
                "Permissions",
                "Hidden",
                "Read-only",
                "Executable",
                "System (Windows)",
                "Archive attribute (Windows)",
                "Temporary (Windows)",
                "Offline (Windows)",
                "Compressed (Windows)",
                "Encrypted (Windows)",
                "Sparse (Windows)",
                "Content indexing disabled (Windows)",
                "File attributes (Windows)",
                "File attributes (st_file_attributes)",
                "Reparse tag (st_reparse_tag)",
                "Mode (st_mode)",
                "File ID (st_ino)",
                "Device ID (st_dev)",
                "Device type (st_rdev)",
                "Link count (st_nlink)",
                "Owner ID (st_uid)",
                "Group ID (st_gid)",
                "Block size (st_blksize)",
                "Allocated blocks (st_blocks)",
                "File size (MiB)",
                "Accessed at",
                "Modified at",
                "Created at",
                "Metadata changed at",
                "BOM",
                "Integrity status",
            }
            self.assertTrue(requested_headers.issubset(columns))
            self.assertEqual(values[columns["Name without extension"]], "leaf")
            self.assertEqual(values[columns["File name"]], "leaf.pdf")
            self.assertEqual(values[columns["Detected content type"]], "PDF")
            self.assertEqual(values[columns["Inside archive"]], "No")
            self.assertEqual(values[columns["BOM"]], "None")
            self.assertIsNone(values[columns["Integrity status"]])
            self.assertEqual(values[columns["Item type"]], "File")
            self.assertEqual(values[columns["Hidden"]], "No")
            self.assertEqual(values[columns["Permissions"]], properties.permissions)
            self.assertEqual(values[columns["Mode (st_mode)"]], properties.mode)
            self.assertEqual(values[columns["File ID (st_ino)"]], properties.inode)
            self.assertIsNotNone(values[columns["Modified at"]])
            status_sheet = workbook["scan_status"]
            status_values = {
                row[0]: row[1]
                for row in status_sheet.iter_rows(
                    min_row=2,
                    max_row=17,
                    values_only=True,
                )
            }
            self.assertEqual(status_values["Status"], "COMPLETE")
            self.assertEqual(status_values["Report rows"], 1)
        finally:
            workbook.close()

    def test_physical_properties_and_utf8_bom(self) -> None:
        source = self.root / "bom.txt"
        source.write_bytes(b"\xef\xbb\xbfhello")

        files, stats = folder_review.scan_tree(str(self.root), set(), workers=2)
        rows = folder_review.inspect_files(files, workers=2)

        self.assertEqual(stats.skipped_entries, 0)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].bom, "UTF-8")
        self.assertEqual(rows[0].content_type, "Text (UTF-8 BOM)")
        self.assertIsNotNone(rows[0].properties)
        self.assertEqual(rows[0].item_type, "File")
        self.assertIsNone(rows[0].link_type)
        self.assertEqual(rows[0].properties.mode, source.stat().st_mode)
        self.assertTrue(rows[0].properties.permissions.startswith("-"))
        if os.name == "nt":
            self.assertIsNotNone(rows[0].properties.file_attributes_value)
            self.assertIsNone(rows[0].properties.metadata_changed_at)
        else:
            self.assertIsNone(rows[0].properties.file_attributes_value)
            self.assertIsNotNone(rows[0].properties.metadata_changed_at)

    def test_folder_enumeration_error_keeps_path_diagnostic(self) -> None:
        with mock.patch.object(
            folder_review.os,
            "scandir",
            side_effect=PermissionError("denied"),
        ):
            subfolders, files, diagnostics, skipped = folder_review.scan_folder(
                str(self.root),
                (),
                set(),
            )

        self.assertEqual(subfolders, [])
        self.assertEqual(len(files), 1)
        self.assertTrue(files[0].virtual_folder)
        self.assertTrue(files[0].folder_unreadable)
        self.assertTrue(skipped)
        self.assertEqual(len(diagnostics), 1)
        self.assertEqual(diagnostics[0].path, str(self.root))
        self.assertEqual(diagnostics[0].stage, "folder enumeration")
        self.assertIn("PermissionError", diagnostics[0].error)

    def test_links_and_special_items_are_reported_without_following(self) -> None:
        if os.name == "nt":
            self.skipTest("POSIX link and FIFO behavior")

        target_file = self.root / "target.txt"
        target_file.write_text("target", encoding="utf-8")
        target_folder = self.root / "target-folder"
        target_folder.mkdir()
        (target_folder / "inside.txt").write_text("inside", encoding="utf-8")
        file_link = self.root / "file-link"
        folder_link = self.root / "folder-link"
        file_link.symlink_to(target_file.name)
        folder_link.symlink_to(target_folder.name, target_is_directory=True)
        fifo = self.root / "events.fifo"
        os.mkfifo(fifo)

        files, stats = folder_review.scan_tree(str(self.root), set(), workers=2)
        rows = folder_review.inspect_files(files, workers=2)
        by_name = {row.leaf_name: row for row in rows}

        self.assertEqual(stats.skipped_entries, 0)
        self.assertEqual(by_name["file-link"].item_type, "Symbolic link")
        self.assertEqual(by_name["file-link"].link_type, "Symbolic link")
        self.assertEqual(by_name["file-link"].link_target, target_file.name)
        self.assertEqual(
            by_name["folder-link"].content_type,
            "Filesystem link (not followed)",
        )
        self.assertEqual(by_name["events.fifo"].item_type, "Named pipe (FIFO)")
        self.assertEqual(
            by_name["events.fifo"].content_type,
            "Special filesystem item (not read)",
        )
        self.assertEqual(sum(row.leaf_name == "inside.txt" for row in rows), 1)

    def test_windows_junction_is_reported_without_following(self) -> None:
        if os.name != "nt":
            self.skipTest("Windows junction behavior")

        target_folder = self.root / "target-folder"
        target_folder.mkdir()
        (target_folder / "inside.txt").write_text("inside", encoding="utf-8")
        junction = self.root / "folder-junction"
        result = subprocess.run(
            ["cmd.exe", "/d", "/c", "mklink", "/J", str(junction), str(target_folder)],
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode != 0:
            self.skipTest(f"Junction creation failed: {result.stderr}")

        files, stats = folder_review.scan_tree(str(self.root), set(), workers=2)
        rows = folder_review.inspect_files(files, workers=2)
        by_name = {row.leaf_name: row for row in rows}

        self.assertEqual(stats.skipped_entries, 0)
        self.assertEqual(
            by_name["folder-junction"].item_type,
            "Directory junction",
        )
        self.assertEqual(by_name["folder-junction"].link_type, "Junction")
        self.assertTrue(by_name["folder-junction"].link_target)
        self.assertEqual(sum(row.leaf_name == "inside.txt" for row in rows), 1)

    def test_hard_links_are_marked(self) -> None:
        source = self.root / "source.txt"
        hard_link = self.root / "hard-link.txt"
        source.write_text("same data", encoding="utf-8")
        try:
            os.link(source, hard_link)
        except OSError as exc:
            self.skipTest(f"Hard links are unavailable: {exc}")

        files, _ = folder_review.scan_tree(str(self.root), set(), workers=2)
        rows = folder_review.inspect_files(files, workers=2)

        self.assertEqual(len(rows), 2)
        self.assertTrue(all(row.link_type == "Hard link" for row in rows))
        self.assertTrue(all(row.properties.links == 2 for row in rows))

    def test_windows_attribute_names(self) -> None:
        self.assertEqual(
            folder_review.windows_attribute_names(0x00000001 | 0x00000020),
            "READONLY | ARCHIVE",
        )
        self.assertEqual(
            folder_review.windows_attribute_names(0x01000000),
            "UNKNOWN_0x01000000",
        )
        self.assertIsNone(folder_review.windows_attribute_names(None))

    def test_common_windows_file_roles(self) -> None:
        cases = {
            "Thumbs.db": "Windows thumbnail cache (Thumbs.db)",
            "ehthumbs.db": "Windows Media Center thumbnail cache (ehthumbs.db)",
            "thumbcache_256.db": "Windows thumbnail cache database",
            "iconcache_32.db": "Windows icon cache database",
            "desktop.ini": "Windows folder settings (desktop.ini)",
            "autorun.inf": "Windows AutoRun configuration (autorun.inf)",
            "~$Quarterly report.xlsx": "Microsoft Office owner/lock file",
            "~Quarterly report.xlsx": "Probable Microsoft Office temporary file",
            "document.lnk": "Windows shortcut",
            "website.url": "Windows Internet shortcut",
            "working.tmp": "Temporary file",
            "recovery.asd": "Microsoft Word AutoRecover file",
            "backup.wbk": "Microsoft Word backup file",
            "backup.xlk": "Microsoft Excel backup file",
        }
        for name, expected in cases.items():
            with self.subTest(name=name):
                self.assertEqual(
                    folder_review.recognized_file_role(name),
                    expected,
                )

        self.assertIsNone(folder_review.recognized_file_role("report.xlsx"))
        self.assertEqual(
            folder_review.recognized_file_role(
                "renamed.bin",
                "Windows shortcut (LNK)",
            ),
            "Windows shortcut",
        )

    def test_windows_objects_keep_content_and_role_separate(self) -> None:
        ole_header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(512)
        lnk_header = (
            b"\x4c\x00\x00\x00\x01\x14\x02\x00\x00\x00"
            b"\x00\x00\xc0\x00\x00\x00\x00\x00\x00\x46" + bytes(32)
        )
        archive = self.make_zip(
            self.root / "windows-objects.zip",
            {
                "Folder/Thumbs.db": ole_header,
                "Folder/~$Budget.xlsx": b"\x00owner",
                "Folder/~Budget.xlsx": b"\x00temporary",
                "Folder/desktop.ini": b"[.ShellClassInfo]\r\nIconFile=folder.ico\r\n",
                "Folder/shortcut.lnk": lnk_header,
            },
        )

        rows = self.inspect(archive)
        by_name = {row.leaf_name: row for row in rows}

        self.assertEqual(
            by_name["Thumbs.db"].content_type,
            "Windows thumbnail cache (Thumbs.db / OLE2)",
        )
        self.assertEqual(
            by_name["Thumbs.db"].recognized_role,
            "Windows thumbnail cache (Thumbs.db)",
        )
        self.assertEqual(by_name["~$Budget.xlsx"].content_type, "Unknown binary")
        self.assertEqual(
            by_name["~$Budget.xlsx"].recognized_role,
            "Microsoft Office owner/lock file",
        )
        self.assertEqual(by_name["~Budget.xlsx"].content_type, "Unknown binary")
        self.assertEqual(
            by_name["~Budget.xlsx"].recognized_role,
            "Probable Microsoft Office temporary file",
        )
        self.assertEqual(
            by_name["desktop.ini"].recognized_role,
            "Windows folder settings (desktop.ini)",
        )
        self.assertEqual(
            by_name["shortcut.lnk"].content_type,
            "Windows shortcut (LNK)",
        )
        self.assertEqual(
            by_name["shortcut.lnk"].recognized_role,
            "Windows shortcut",
        )
        self.assertTrue(all(row.in_archive == "Yes" for row in rows))

    def test_all_supported_bom_variants(self) -> None:
        cases = {
            b"\xef\xbb\xbftext": "UTF-8",
            b"\xff\xfe\x00\x00text": "UTF-32 LE",
            b"\x00\x00\xfe\xfftext": "UTF-32 BE",
            b"\xff\xfetext": "UTF-16 LE",
            b"\xfe\xfftext": "UTF-16 BE",
            b"plain": None,
        }
        for data, expected in cases.items():
            with self.subTest(expected=expected):
                self.assertEqual(folder_review.detect_bom(data), expected)

    def test_platform_and_worker_defaults(self) -> None:
        self.assertEqual(folder_review.DEFAULT_WORKERS, 16)
        if os.name == "nt":
            self.assertEqual(
                folder_review.regular_path(r"\\?\UNC\server\share\folder"),
                r"\\server\share\folder",
            )
            self.assertEqual(
                folder_review.regular_path(r"\\?\unc\server\share\folder"),
                r"\\server\share\folder",
            )
        with (
            mock.patch.object(folder_review.sys, "platform", "darwin"),
            self.assertRaisesRegex(RuntimeError, "macOS is not supported"),
        ):
            folder_review.validate_platform()

    def test_partial_scan_writes_diagnostics_and_returns_two(self) -> None:
        input_root = self.root / "input"
        input_root.mkdir()
        output = self.root / "partial.xlsx"
        diagnostic = folder_review.Diagnostic(
            str(input_root / "unavailable"),
            "folder enumeration",
            "PermissionError: denied",
        )
        stats = folder_review.ScanStats(
            folders=2,
            skipped_folders=1,
            diagnostics=[diagnostic],
        )
        with (
            mock.patch.object(
                folder_review.sys,
                "argv",
                [
                    "folder_review.py",
                    "--root",
                    str(input_root),
                    "--output",
                    str(output),
                ],
            ),
            mock.patch.object(folder_review, "scan_tree", return_value=([], stats)),
            contextlib.redirect_stdout(io.StringIO()),
            contextlib.redirect_stderr(io.StringIO()),
        ):
            exit_code = folder_review.run()

        self.assertEqual(exit_code, 2)
        workbook = load_workbook(output, read_only=True, data_only=True)
        try:
            status_sheet = workbook["scan_status"]
            values = list(status_sheet.iter_rows(values_only=True))
            self.assertIn(("Status", "PARTIAL"), [row[:2] for row in values])
            self.assertTrue(
                any(
                    len(row) >= 3
                    and row[0] == "folder enumeration"
                    and row[1] == "unavailable"
                    and "PermissionError" in row[2]
                    for row in values
                )
            )
        finally:
            workbook.close()

    def test_physical_row_limit_writes_partial_workbook(self) -> None:
        input_root = self.root / "input"
        input_root.mkdir()
        for number in range(3):
            (input_root / f"file-{number}.txt").write_text(
                "plain text",
                encoding="utf-8",
            )
        output = self.root / "row-limit.xlsx"
        with (
            mock.patch.object(
                folder_review.sys,
                "argv",
                [
                    "folder_review.py",
                    "--root",
                    str(input_root),
                    "--output",
                    str(output),
                    "--max-rows",
                    "2",
                ],
            ),
            contextlib.redirect_stdout(io.StringIO()),
            contextlib.redirect_stderr(io.StringIO()),
        ):
            exit_code = folder_review.run()

        self.assertEqual(exit_code, 2)
        self.assertTrue(output.is_file())
        workbook = load_workbook(output, read_only=True, data_only=True)
        try:
            inventory = workbook["folder_review"]
            self.assertEqual(
                sum(1 for _ in inventory.iter_rows(min_row=7, values_only=True)),
                2,
            )
            status_values = list(workbook["scan_status"].iter_rows(values_only=True))
            self.assertIn(("Status", "PARTIAL"), [row[:2] for row in status_values])
            self.assertIn(("Physical items", 3), [row[:2] for row in status_values])
            self.assertIn(("Report rows", 2), [row[:2] for row in status_values])
            self.assertTrue(
                any(
                    len(row) >= 3
                    and row[0] == "row limit"
                    and row[1] == "."
                    and "1 physical items omitted" in row[2]
                    for row in status_values
                )
            )
        finally:
            workbook.close()

    def test_auto_workers_are_bounded_and_manual_override_is_preserved(self) -> None:
        getter = (
            "process_cpu_count"
            if hasattr(folder_review.os, "process_cpu_count")
            else "cpu_count"
        )
        for cpus, remote, expected in (
            (1, False, 4),
            (4, False, 8),
            (64, False, 16),
            (1, True, 8),
            (4, True, 16),
            (64, True, 32),
            (None, False, 4),
        ):
            with (
                mock.patch.object(folder_review.os, getter, return_value=cpus),
                mock.patch.object(
                    folder_review, "local_storage", return_value=not remote
                ),
            ):
                self.assertEqual(folder_review.choose_workers("input"), expected)
        self.assertIsNone(folder_review.worker_count("auto"))
        self.assertEqual(folder_review.worker_count("23"), 23)
        with mock.patch.object(
            folder_review.sys, "argv", ["folder_review.py", "--workers", "23"]
        ):
            self.assertEqual(folder_review.parse_args().workers, 23)

    def test_pasted_paths_preserve_unc_spaces_unicode_and_matching_quotes(self) -> None:
        for raw, expected in (
            (' "C:\\Folder with spaces\\пример" ', "C:\\Folder with spaces\\пример"),
            ("'\\\\server\\share\\folder'", "\\\\server\\share\\folder"),
            ("/data/folder\r\n", "/data/folder"),
        ):
            self.assertEqual(folder_review.pasted_path(raw), expected)

    def test_workbook_freeze_widths_and_dax_reference_are_text(self) -> None:
        output = self.root / "layout.xlsx"
        folder_review.write_workbook([], output, "sample")
        book = load_workbook(output)
        try:
            self.assertEqual(
                book.sheetnames, ["folder_review", "scan_status", "dax_reference"]
            )
            self.assertEqual(book["folder_review"].freeze_panes, "A7")
            self.assertEqual(book["folder_review"].sheet_view.zoomScale, 70)
            self.assertGreater(book["scan_status"].column_dimensions["A"].width, 30)
            self.assertGreater(book["scan_status"].column_dimensions["B"].width, 60)
            for sheet in (book["folder_review"], book["dax_reference"]):
                self.assertTrue(
                    all(
                        abs(dim.width - 5.7109375) < 0.01
                        for dim in sheet.column_dimensions.values()
                    )
                )
            cell = book["dax_reference"]["A1"]
            self.assertEqual(cell.value, folder_review.DAX_MEASURE)
            self.assertEqual(cell.data_type, "s")
            self.assertEqual(cell.value.count("'Data'[TextColumn]"), 1)
            self.assertTrue(cell.value.startswith("=VAR Source"))
        finally:
            book.close()

    def test_relative_report_paths_preserve_archive_members_and_source_names(self) -> None:
        cases = (
            ("/data/пример/文件.txt", "/data/пример", "文件.txt"),
            ("/data/a.zip::папка/b.7z::文件.txt", "/data", "a.zip::папка/b.7z::文件.txt"),
            ("/data/empty", "/data", "empty"),
            ("/data", "/data", "."),
            ("/file.txt", "/", "file.txt"),
            (r"C:\Data\папка\文件.txt", r"c:\data", r"папка\文件.txt"),
            (r"C:\Data\a.zip::папка/b.txt", r"C:\Data", "a.zip::папка/b.txt"),
            ("C:\\file.txt", "C:\\", "file.txt"),
            (r"\\server\share\папка\文件.txt", r"\\server\share", r"папка\文件.txt"),
            (r"D:\external.txt", r"C:\Data", r"D:\external.txt"),
        )
        for path, root, expected in cases:
            with self.subTest(path=path, root=root):
                self.assertEqual(folder_review.report_relative_path(path, root), expected)

    def test_report_technical_values_are_english_and_root_is_only_in_title(self) -> None:
        root = self.root / "пример"
        root.mkdir()
        source = root / "文件.txt"
        source.write_bytes(b"content")
        rows = [replace(folder_review.make_row(
            str(source), (source.name,), 7, "Text", False,
            properties=folder_review.file_properties(source.name, source.stat()),
            bom=folder_review.BOM_NOT_PRESENT,
        ), content_hash=blake3(b"content").digest()),
            folder_review.folder_placeholder(str(root / "пусто"), ("пусто",), in_archive=False)]
        output = self.root / "english.xlsx"
        folder_review.write_workbook(rows, output, str(root), full_metadata=True)
        book = load_workbook(output)
        try:
            sheet = book["folder_review"]
            self.assertEqual(sheet["B2"].value, f"Folder file inventory: {root}")
            self.assertEqual(sum(str(root) in str(cell.value) for row in sheet for cell in row), 1)
            headers = [cell.value for cell in sheet[6]]
            records = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=7, values_only=True)]
            self.assertNotIn("Full path", headers)
            self.assertEqual(records[0]["Relative path"], "文件.txt")
            self.assertEqual(records[0]["File name"], "文件.txt")
            self.assertEqual(records[1]["Relative path"], "пусто")
            self.assertEqual(records[1]["Inspection state"], "No inventoried files")
            source_columns = {None, "Relative path", "File name", "Name without extension", "Level 1"}
            for record in records:
                for header, value in record.items():
                    if header not in source_columns and isinstance(value, str):
                        self.assertTrue(value.isascii(), (header, value))
            self.assertEqual(folder_review.INTEGRITY_VERIFIED, "Verified")
            self.assertEqual(folder_review.INTEGRITY_NOT_VERIFIED, "Not verified")
            self.assertEqual(folder_review.INTEGRITY_ERROR, "Failed")
        finally:
            book.close()

    def test_in_folder_report_is_excluded_on_repeated_scans(self) -> None:
        root = self.root / "target"
        root.mkdir()
        source = root / "file.txt"
        source.write_bytes(b"unchanged")
        output = root / "folder-review.xlsx"
        for _ in range(2):
            with (
                mock.patch.object(
                    folder_review.sys,
                    "argv",
                    ["folder_review.py", "--root", str(root), "--output", str(output)],
                ),
                contextlib.redirect_stdout(io.StringIO()),
            ):
                self.assertEqual(folder_review.run(), 0)
            book = load_workbook(output, read_only=True)
            try:
                sheet = book["folder_review"]
                headers = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
                records = list(sheet.iter_rows(min_row=7, values_only=True))
                self.assertEqual(
                    [row[headers.index("File name")] for row in records], ["file.txt"]
                )
            finally:
                book.close()
        self.assertEqual(source.read_bytes(), b"unchanged")
        self.assertEqual(set(root.iterdir()), {source, output})

    def test_tool_in_target_uses_default_report_without_an_extra_prompt(self) -> None:
        root = self.root / "tool folder"
        root.mkdir()
        tool = root / "tool.py"
        tool.write_bytes(b"ordinary source file")
        with (
            mock.patch.object(folder_review, "__file__", str(tool)),
            mock.patch.object(
                folder_review.sys, "argv", [str(tool), "--root", str(root)]
            ),
            mock.patch(
                "builtins.input", side_effect=AssertionError("Unexpected prompt")
            ),
            contextlib.redirect_stdout(io.StringIO()),
        ):
            self.assertEqual(folder_review.run(), 0)
        self.assertTrue((root / "folder-review.xlsx").is_file())
        self.assertEqual(tool.read_bytes(), b"ordinary source file")

    def test_interactive_first_run_prompts_for_folder(self) -> None:
        input_root = self.root / "input"
        input_root.mkdir()
        output = self.root / "prompted.xlsx"
        with (
            mock.patch.object(
                folder_review.sys,
                "argv",
                ["folder_review.py", "--output", str(output)],
            ),
            mock.patch.object(folder_review.sys.stdin, "isatty", return_value=True),
            mock.patch("builtins.input", return_value=str(input_root)),
            contextlib.redirect_stdout(io.StringIO()),
        ):
            exit_code = folder_review.run()

        self.assertEqual(exit_code, 0)
        self.assertTrue(output.is_file())

    def test_archive_path_splits_both_separator_styles(self) -> None:
        self.assertEqual(
            folder_review.split_archive_path(r"First/Second\Third/file.txt"),
            ("First", "Second", "Third", "file.txt"),
        )

    def test_default_columns_match_full_metadata_values(self) -> None:
        source = self.root / ".hidden.pdf"
        source.write_bytes(b"%PDF-1.4\n")
        rows = self.inspect(source)
        reports = []
        for full in (False, True):
            output = self.root / f"metadata-{full}.xlsx"
            folder_review.write_workbook(rows, output, "sample", full_metadata=full)
            workbook = load_workbook(output, read_only=True, data_only=False)
            try:
                sheet = workbook["folder_review"]
                headers = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
                cells = next(sheet.iter_rows(min_row=7, max_row=7))
                reports.append(
                    {
                        header: (cell.value, cell.data_type, cell.number_format)
                        for header, cell in zip(headers, cells)
                        if header
                    }
                )
            finally:
                workbook.close()
        compact, full = reports
        self.assertEqual(
            set(full) - set(compact), folder_review.ADVANCED_METADATA_HEADERS
        )
        self.assertIn("Hidden", compact)
        self.assertEqual(compact, {key: full[key] for key in compact})

    def test_runtime_denies_network_and_source_mutations(self) -> None:
        import socket

        input_root = self.root / "input"
        input_root.mkdir()
        source = input_root / "data.txt"
        source.write_bytes(b"unchanged")
        before = source.stat()
        output = self.root / "report.xlsx"
        with folder_review.safe_runtime(str(input_root), output):
            null_descriptor = os.open(os.devnull, os.O_RDWR)
            os.close(null_descriptor)
            operations = [
                lambda: source.write_bytes(b"changed"),
                source.unlink,
                lambda: source.rename(self.root / "moved.txt"),
                lambda: (input_root / "new").mkdir(),
                lambda: source.chmod(0o600),
                lambda: os.utime(source, (1, 1)),
                lambda: os.truncate(source, 0),
                lambda: os.link(source, self.root / "hard-link.txt"),
                lambda: socket.socket(),
                lambda: socket.getaddrinfo("example.invalid", 443),
                lambda: subprocess.run(
                    [folder_review.sys.executable, "--version"], check=True
                ),
                lambda: subprocess.run(["7z", "e", str(source)], check=True),
            ]
            for operation in operations:
                with (
                    self.subTest(operation=operations.index(operation)),
                    self.assertRaises(PermissionError),
                ):
                    operation()
            output.write_bytes(b"allowed outside input")
            with mock.patch.object(folder_review.shutil, "which", return_value=None):
                self.assertFalse(
                    folder_review.rarfile.ToolSetup(
                        folder_review.rarfile.UNRAR_CONFIG
                    ).check()
                )
            with (
                mock.patch.object(
                    folder_review.shutil, "which", return_value=str(source)
                ),
                self.assertRaisesRegex(PermissionError, "read-only"),
            ):
                subprocess.run(["7z", "i"], check=True)
            with (
                mock.patch.object(folder_review, "local_storage", return_value=False),
                self.assertRaisesRegex(PermissionError, "local storage"),
            ):
                (self.root / "remote-work.tmp").write_bytes(b"forbidden")
        self.assertEqual(source.read_bytes(), b"unchanged")
        self.assertEqual(source.stat().st_mtime_ns, before.st_mtime_ns)
        self.assertEqual(source.stat().st_mode, before.st_mode)
        self.assertEqual(list(input_root.iterdir()), [source])

    def test_only_selected_report_can_be_written_inside_input_and_aliases(self) -> None:
        input_root = self.root / "input"
        input_root.mkdir()
        output = input_root / "report.xlsx"
        with folder_review.safe_runtime(str(input_root), output):
            with self.assertRaises(PermissionError):
                output.write_bytes(b"not a report save")
            folder_review.write_workbook([], output, str(input_root))
            with self.assertRaises(PermissionError):
                folder_review.write_workbook(
                    [], input_root / "other.xlsx", str(input_root)
                )
            with self.assertRaises(PermissionError):
                (input_root / "other.txt").write_bytes(b"forbidden")
        self.assertTrue(output.is_file())
        output.unlink()
        alias = self.root / "alias"
        if os.name == "nt":
            subprocess.run(
                ["cmd.exe", "/c", "mklink", "/J", str(alias), str(input_root)],
                check=True,
                capture_output=True,
            )
        else:
            alias.symlink_to(input_root, target_is_directory=True)
        try:
            folder_review.validate_output_path(alias / "report.xlsx", str(input_root))
            with folder_review.safe_runtime(str(input_root), alias / "report.xlsx"):
                folder_review.write_workbook([], alias / "report.xlsx", str(input_root))
            self.assertIn(
                folder_review.normalized_key(output),
                folder_review.report_exclusions(alias / "report.xlsx", str(input_root)),
            )
            output.unlink()
            with (
                folder_review.safe_runtime(str(input_root), self.root / "outside.xlsx"),
                self.assertRaises(PermissionError),
            ):
                (alias / "changed.txt").write_bytes(b"forbidden")
        finally:
            if os.name == "nt":
                alias.rmdir()
            else:
                alias.unlink()
        self.assertEqual(list(input_root.iterdir()), [])

    def test_temporary_files_stay_outside_input_and_are_cleaned(self) -> None:
        input_root = self.root / "input"
        input_root.mkdir()
        original_tempdir = tempfile.tempdir
        with mock.patch.object(tempfile, "tempdir", str(input_root)):
            with folder_review.safe_runtime(str(input_root), self.root / "output.xlsx"):
                workspace = Path(tempfile.gettempdir())
                self.assertFalse(folder_review.path_within(workspace, str(input_root)))
                with tempfile.TemporaryFile() as handle:
                    handle.write(b"local work")
                with folder_review.archive_temp_directory() as directory:
                    (Path(directory) / "work.bin").write_bytes(b"temporary")
            self.assertFalse(workspace.exists())
            self.assertEqual(tempfile.tempdir, str(input_root))
        self.assertEqual(tempfile.tempdir, original_tempdir)
        self.assertEqual(list(input_root.iterdir()), [])

    def test_temp_directory_discovery_cannot_write_to_input(self) -> None:
        input_root = self.root / "input"
        input_root.mkdir()
        probe = input_root / "probe.tmp"
        with (
            mock.patch.object(
                tempfile, "gettempdir", side_effect=lambda: probe.write_bytes(b"probe")
            ),
            self.assertRaises(PermissionError),
            folder_review.safe_runtime(str(input_root), self.root / "output.xlsx"),
        ):
            self.fail("Input was allowed as a writable temp candidate")
        self.assertFalse(probe.exists())
        self.assertIsNone(folder_review._protected_root)
        with (
            mock.patch.object(tempfile, "tempdir", None),
            mock.patch.dict(
                os.environ,
                {name: str(input_root) for name in ("TMPDIR", "TEMP", "TMP")},
            ),
            folder_review.safe_runtime(str(input_root), self.root / "output.xlsx"),
        ):
            self.assertFalse(
                folder_review.path_within(tempfile.gettempdir(), str(input_root))
            )
        self.assertEqual(list(input_root.iterdir()), [])

    def test_long_physical_archive_and_output_paths(self) -> None:
        deep_folder = str(self.root)
        segment = 0
        while len(deep_folder) < 320:
            deep_folder = os.path.join(
                deep_folder,
                f"long-folder-segment-{segment:02d}-abcdefgh",
            )
            segment += 1
        os.makedirs(folder_review.extended_path(deep_folder))

        inner = self.make_zip(
            self.root / "long-inner.zip",
            {"Inside/utf8-bom.txt": b"\xef\xbb\xbflong path"},
        )
        outer = self.make_7z(
            self.root / "long-outer.7z",
            {"Archive folder/long-inner.zip": inner},
        )
        long_archive = os.path.join(deep_folder, "long-outer.7z")
        with (
            open(outer, "rb") as source,
            open(folder_review.extended_path(long_archive), "wb") as destination,
        ):
            destination.write(source.read())

        files, stats = folder_review.scan_tree(deep_folder, set(), workers=2)
        rows = folder_review.inspect_files(files, workers=2)

        self.assertEqual(stats.skipped_folders, 0)
        self.assertEqual(stats.skipped_entries, 0)
        self.assertEqual(len(files), 1)
        self.assertEqual(len(rows), 1)
        self.assertGreater(len(rows[0].full_path), 260)
        self.assertEqual(rows[0].leaf_name, "utf8-bom.txt")
        self.assertEqual(rows[0].bom, "UTF-8")
        self.assertIsNone(rows[0].error)

        output_folder = os.path.join(deep_folder, "workbook-output")
        os.mkdir(folder_review.extended_path(output_folder))
        output = Path(output_folder) / "folder-review.xlsx"
        folder_review.write_workbook(rows, output, deep_folder)

        self.assertTrue(os.path.isfile(folder_review.extended_path(output)))
        self.assertFalse(
            any(
                name.startswith(".folder-review-")
                for name in os.listdir(folder_review.extended_path(output_folder))
            )
        )
        workbook = load_workbook(
            folder_review.extended_path(output),
            read_only=True,
            data_only=False,
        )
        try:
            sheet = workbook["folder_review"]
            values = next(sheet.iter_rows(min_row=7, max_row=7, values_only=True))
            self.assertIn("UTF-8", values)
            self.assertTrue(
                any(isinstance(value, str) and len(value) > 260 for value in [sheet["B2"].value, *values])
            )
        finally:
            workbook.close()

    def test_temp_member_path_uses_safe_short_random_name(self) -> None:
        with tempfile.TemporaryDirectory(prefix="folder-review-temp-test-") as root:
            path = folder_review.make_temp_path(
                root,
                "member-with-invalid-extension." + "x" * 300 + "?zip",
            )
            try:
                self.assertLess(len(os.path.basename(path)), 80)
                self.assertTrue(path.endswith(".bin"))
            finally:
                os.unlink(path)


if __name__ == "__main__":
    unittest.main()
