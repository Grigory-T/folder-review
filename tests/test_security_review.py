"""Regressions for unusual names, mount scope and owned report staging."""

import contextlib
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import folder_change_review as changes
import folder_review as core
import review_state as states


class SecurityReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.tool = self.base / "tool"
        self.tool.mkdir()
        self.root = self.base / "input"
        self.root.mkdir()
        location = patch.object(core, "application_directory", return_value=self.tool)
        location.start()
        self.addCleanup(location.stop)

    def capture(self, name):
        output = self.tool / name
        with contextlib.redirect_stdout(io.StringIO()):
            code = core.execute_review(core.parse_args([]), str(self.root), output)
        return output, states.load_snapshot(output), code

    def test_visible_text_escapes_surrogates_and_keeps_normal_unicode(self):
        self.assertEqual(core.excel_text("文😀=x"), "文😀=x")
        self.assertEqual(core.excel_text("bad-\udcff-\ud800"), r"bad-\udcff-\ud800")
        self.assertEqual(core.excel_text("bad\x01\ufffe\uffff"), "bad���")
        self.assertIsNone(core.excel_text(None))
        self.assertEqual(len(core.excel_text("x" * 32768, truncate=False)), 32768)

    def test_unusual_filename_retains_exact_state_and_both_reports_work(self):
        name = (
            "invalid-\udcff.txt"
            if os.name == "nt"
            else os.fsdecode(b"invalid-\xff.txt")
        )
        file = self.root / name
        file.write_bytes(b"unchanged")
        original_time = file.stat().st_mtime_ns
        output, snapshot, code = self.capture("before.xlsx")
        self.assertEqual(code, 0)
        self.assertEqual(snapshot["files"][0]["parts"], [name])
        with contextlib.closing(load_workbook(output)) as book:
            values = list(book["folder_review"].iter_rows(min_row=6, values_only=True))
            self.assertEqual(
                values[1][values[0].index("File name")], r"invalid-\udcff.txt"
            )
        after, _, code = self.capture("after.xlsx")
        self.assertEqual(code, 0)
        result = self.tool / "changes.xlsx"
        with contextlib.redirect_stdout(io.StringIO()):
            counts = changes.compare_files(output, after, result)
        self.assertEqual(counts, {"Unchanged": 1})
        with contextlib.closing(load_workbook(result)) as book:
            self.assertIn(r"invalid-\udcff.txt", [c.value for c in book["Changes"][7]])
        self.assertEqual(file.read_bytes(), b"unchanged")
        self.assertEqual(file.stat().st_mtime_ns, original_time)

    def test_comparison_roots_and_read_issues_are_valid_literal_xml_text(self):
        output = self.tool / "partial.xlsx"
        root = str(self.root) + "\ud800"
        issue = {
            "path": root + os.sep + "=bad\udcff",
            "kind": "file",
            "errno": 13,
            "winerror": None,
            "message": "Denied\ud800\ufffe",
        }
        with core.safe_runtime(str(self.root), output):
            changes.write_report(
                [],
                output,
                root_hint=root,
                read_errors=[("A", issue)],
                issue_roots={"A": root},
            )
        with contextlib.closing(load_workbook(output)) as book:
            self.assertIn(r"\ud800", book["Changes"]["B2"].value)
            self.assertIn(r"\ud800", book["Read issues"]["B2"].value)
            self.assertEqual(book["Read issues"]["C6"].value, r"=bad\udcff")
            self.assertEqual(book["Read issues"]["C6"].data_type, "s")
            self.assertEqual(book["Read issues"]["E6"].value, "Denied\\ud800�")

    @unittest.skipIf(os.name == "nt", "Linux mount metadata")
    def test_mount_snapshot_includes_bind_mounts_and_decodes_paths(self):
        mount = str(self.root / "same device")
        escaped = mount.replace(" ", r"\040")
        text = (
            "1 0 8:1 / / rw - ext4 /dev/root rw\n"
            f"2 1 8:1 /other {escaped} rw - ext4 /dev/root rw\n"
        )
        with patch.object(core, "open", return_value=io.StringIO(text), create=True):
            self.assertEqual(core.nested_mounts(str(self.root)), frozenset({mount}))

    def test_submount_is_not_touched_and_missing_coverage_is_unknown(self):
        blocked = self.root / "mounted"
        blocked.mkdir()
        (blocked / "leaf.txt").write_bytes(b"outside scope")
        before, _, _ = self.capture("before.xlsx")
        actual_scandir = os.scandir

        class Entry:
            def __init__(self, original):
                self.name = original.name

            def stat(self, **kwargs):
                raise AssertionError("Mount target metadata was accessed")

        class Scan:
            def __init__(self, path):
                if not isinstance(path, int):
                    self_path = core.regular_path(os.fsdecode(path))
                    if self_path == str(blocked):
                        raise AssertionError("Mount target was enumerated")
                self.entries = actual_scandir(path)

            def __iter__(self):
                return self

            def __next__(self):
                entry = next(self.entries)
                return Entry(entry) if entry.name == "mounted" else entry

            def __enter__(self):
                return self

            def __exit__(self, *args):
                self.close()

            def close(self):
                self.entries.close()

        with (
            patch.object(core, "nested_mounts", return_value=frozenset({str(blocked)})),
            patch.object(core.os, "scandir", side_effect=Scan),
        ):
            after, snapshot, code = self.capture("after.xlsx")
        self.assertEqual(code, 2)
        self.assertFalse(snapshot["complete"])
        self.assertEqual(snapshot["read_errors"][0]["kind"], "directory")
        with contextlib.redirect_stdout(io.StringIO()):
            counts = changes.compare_files(before, after, self.tool / "changes.xlsx")
        self.assertEqual(counts, {"Unknown": 1})
        self.assertTrue(counts.partial)

    def test_staging_is_owned_and_legacy_crash_files_are_excluded_not_deleted(self):
        output = self.tool / "report.xlsx"
        stale = self.tool / (".folder-review-" + "a" * 32 + ".tmp.xlsx")
        stale.write_bytes(b"old partial report")
        unrelated = self.root / stale.name
        unrelated.write_bytes(b"unrelated source")
        with (
            core.safe_runtime(str(self.root), output),
            core._report_workbook(output) as book,
        ):
            book.add_worksheet("data").write_string(0, 0, "ok")
            workspace = Path(tempfile.gettempdir())
            self.assertTrue(core.path_within(workspace, str(core.portable_directory())))
            self.assertEqual(len(list(workspace.glob(".folder-review-*.tmp.xlsx"))), 1)
            self.assertEqual(list(self.tool.glob(".folder-review-*.tmp.xlsx")), [stale])
        self.root = self.tool
        _, snapshot, code = self.capture("state.xlsx")
        self.assertEqual(code, 0)
        self.assertNotIn(stale.name, [r["parts"][-1] for r in snapshot["files"]])
        self.assertEqual(stale.read_bytes(), b"old partial report")
        self.assertEqual(unrelated.read_bytes(), b"unrelated source")
        self.assertFalse(list(core.portable_directory().glob("work-*")))


if __name__ == "__main__":
    unittest.main()
